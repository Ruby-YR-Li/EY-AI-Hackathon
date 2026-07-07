from __future__ import annotations

import csv
import json
import re
import shutil
from copy import copy
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "outputs"
UPLOAD_DIR = BASE_DIR / "uploads"
SAMPLE_DIR = BASE_DIR / "sample"
RULE_DIR = BASE_DIR / "rules"

AUDIT_SHEETS = [
    "汇总",
    "Uexp.00 Lead",
    "VC.00 销售费用BKD",
    "VD.00 管理费用BKD",
    "VC&VD.01.2 详细测试 TOD",
    "VD.01.3 复核法律费用",
    "VC&VD.01.4 截止性测试",
    "VC&VD.01.2 TOD 年审",
    "VC&VD.01.2 TOD 预审+剩余期间",
    "VC&VD.01.2 TOD 抽样工具输出",
]

INTERNAL_PREFIXES = ("Skywind", "DS_INTERNAL")
ERROR_TOKENS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!")
EMPTY_VALUES = (None, "")


@dataclass
class ReviewNote:
    id: str
    category: str
    severity: str
    sheet: str
    cell: str
    issue: str
    audit_risk: str
    basis: str
    suggestion: str
    confidence: float = 1.0
    source: str = "rule"

    @property
    def location(self) -> str:
        return f"{self.sheet}!{self.cell}" if self.cell else self.sheet


def ensure_dirs() -> None:
    for path in (OUTPUT_DIR, UPLOAD_DIR, SAMPLE_DIR, RULE_DIR):
        path.mkdir(exist_ok=True)


def is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def norm(value) -> str:
    return "" if value is None else str(value).strip()


def parse_number(value):
    if isinstance(value, (int, float)):
        return float(value)
    if value is None:
        return None
    text = str(value).replace(",", "").replace("，", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


def is_yes(value) -> bool:
    return norm(value).lower() in {"是", "yes", "y", "true", "执行"}


def is_no(value) -> bool:
    return norm(value).lower() in {"否", "no", "n", "false", "不执行"}


def is_audit_sheet(sheet_name: str) -> bool:
    return not sheet_name.startswith(INTERNAL_PREFIXES)


def add_note(notes: list[ReviewNote], category: str, severity: str, sheet: str, cell: str, issue: str, audit_risk: str, basis: str, suggestion: str) -> None:
    notes.append(
        ReviewNote(
            id=f"RN-{len(notes) + 1:03d}",
            category=category,
            severity=severity,
            sheet=sheet,
            cell=cell,
            issue=issue,
            audit_risk=audit_risk,
            basis=basis,
            suggestion=suggestion,
        )
    )


def load_checkpoints() -> list[dict]:
    path = RULE_DIR / "vcvd_checkpoints.json"
    if not path.exists():
        return []
    return json.loads(path.read_text(encoding="utf-8"))


def scan_workbook(path: Path) -> dict:
    wb = load_workbook(path, data_only=False)
    visible_sheets = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    formulas = 0
    filled = 0
    errors = 0
    for ws in wb.worksheets:
        if not is_audit_sheet(ws.title):
            continue
        for row in ws.iter_rows():
            for cell in row:
                if cell.value not in EMPTY_VALUES:
                    filled += 1
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
                    if any(token in cell.value for token in ERROR_TOKENS):
                        errors += 1
                elif isinstance(cell.value, str) and cell.value.strip() in ERROR_TOKENS:
                    errors += 1
    return {
        "file": path.name,
        "sheet_count": len(wb.sheetnames),
        "visible_sheets": visible_sheets,
        "filled_cells": filled,
        "formula_count": formulas,
        "formula_error_count": errors,
    }


def review_workbook(path: Path) -> tuple[list[ReviewNote], dict]:
    wb = load_workbook(path, data_only=False)
    wb_values = load_workbook(path, data_only=True)
    notes: list[ReviewNote] = []

    check_expected_sheets(wb, notes, path)
    check_formula_errors(wb, notes)
    check_lead_sheet(wb, notes)
    check_summary(wb, notes)
    check_bkd(wb, notes, wb_values)
    check_tod(wb, notes, wb_values)
    check_legal_fee(wb, notes)
    check_cutoff(wb, notes, wb_values)
    check_professional_wording(wb, notes)

    stats = build_stats(notes)
    return notes, stats


def build_stats(notes: list[ReviewNote]) -> dict:
    stats = {
        "total": len(notes),
        "high": sum(1 for n in notes if n.severity == "High"),
        "medium": sum(1 for n in notes if n.severity == "Medium"),
        "low": sum(1 for n in notes if n.severity == "Low"),
        "categories": {},
    }
    for note in notes:
        stats["categories"][note.category] = stats["categories"].get(note.category, 0) + 1
    return stats


def check_expected_sheets(wb: Workbook, notes: list[ReviewNote], source_path: Path | None = None) -> None:
    is_tod_only = any("TOD 年审" in s or "TOD 预审" in s for s in wb.sheetnames) and "汇总" not in wb.sheetnames
    expected = ["VC&VD.01.2 TOD 抽样工具输出"] if is_tod_only else [
        "汇总",
        "Uexp.00 Lead",
        "VC.00 销售费用BKD",
        "VD.00 管理费用BKD",
        "VC&VD.01.4 截止性测试",
    ]
    for sheet in expected:
        if sheet not in wb.sheetnames:
            add_note(
                notes,
                "重点审计程序检查",
                "High",
                sheet,
                "",
                f"标准费用底稿页签缺失：{sheet}。",
                "标准页签缺失可能导致必要 PSP 未执行或 Review 无法追踪对应工作。",
                "SOP 汇总易错点：未使用标准底稿，导致未执行必要 PSP，或 PSP 执行不到位。",
                "请使用标准 VC&VD 费用底稿模板补回该页签，或在汇总页明确说明替代程序及交叉索引。",
            )
    if not is_tod_only and "汇总" in wb.sheetnames:
        ws = wb["汇总"]
        has_external_tod = source_path is not None and any(
            p != source_path and p.suffix.lower() in {".xlsx", ".xlsm"} and "TOD" in p.name.upper() and not p.name.startswith("~$")
            for p in source_path.parent.glob("*.xls*")
        )
        for row in range(3, ws.max_row + 1):
            page = norm(ws[f"F{row}"].value)
            if page in {"程序页", "底稿索引", "N/A"}:
                continue
            if is_yes(ws[f"G{row}"].value) and page and "TOD" in page.upper() and has_external_tod:
                continue
            if is_yes(ws[f"G{row}"].value) and page and page not in wb.sheetnames and "项目组添加底稿索引" not in page:
                add_note(
                    notes,
                    "重点审计程序检查",
                    "High",
                    "汇总",
                    f"F{row}",
                    f"汇总页标记执行“{page}”，但工作簿中未找到对应页签。",
                    "程序被标记执行但缺少对应底稿，可能导致 Manager 无法复核执行过程和审计证据。",
                    "SOP 汇总易错点：实际执行的程序与程序要求不一致，或少执行必要审计程序。",
                    "请补回对应底稿页或补充清晰的外部底稿索引；如未执行，应更新执行状态并说明原因。",
                )


def check_formula_errors(wb: Workbook, notes: list[ReviewNote]) -> None:
    for ws in wb.worksheets:
        if not is_audit_sheet(ws.title):
            continue
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                has_error = isinstance(value, str) and (value.strip() in ERROR_TOKENS or any(token in value for token in ERROR_TOKENS))
                if has_error:
                    add_note(
                        notes,
                        "逻辑与引用检查",
                        "High",
                        ws.title,
                        cell.coordinate,
                        f"发现公式或单元格错误：{value}。",
                        "引用错误会影响合计、阈值或测试范围，可能导致后续审计结论基于错误数据。",
                        "SOP 多处要求金额、阈值与底稿/报表核对一致；公式引用错误属于底稿逻辑缺陷。",
                        "请检查该单元格的引用来源，修复公式后重新刷新底稿并复核相关下游计算。",
                    )


def check_lead_sheet(wb: Workbook, notes: list[ReviewNote]) -> None:
    if "Uexp.00 Lead" not in wb.sheetnames:
        return
    ws = wb["Uexp.00 Lead"]
    fields = [
        ("C2", "客户名称", "Medium"),
        ("C3", "期末", "High"),
        ("C4", "分析日期", "Medium"),
        ("C5", "可容忍误差 TE", "High"),
        ("C6", "名义金额 SAD", "High"),
        ("C7", "适用会计准则", "Medium"),
        ("C8", "记账本位币/单位", "Medium"),
    ]
    for coord, label, severity in fields:
        if is_blank(ws[coord].value):
            add_note(
                notes,
                "基础完整性检查",
                severity,
                ws.title,
                coord,
                f"Lead Sheet 关键字段“{label}”未填写。",
                "基础项缺失会影响底稿识别、阈值计算、期间判断或交付一致性。",
                "SOP Lead Sheet 易错点：PM/TE/SAD 与 Canvas 结果一致；客户名称、期末、货币/单位等基础信息需更新。",
                f"请补充“{label}”，并与 Canvas / 项目基础信息核对一致。",
            )
    standard = norm(ws["C7"].value).upper()
    currency = norm(ws["C8"].value).upper()
    standard_like_currency = standard in {"CNY", "RMB", "USD", "HKD"} or standard in {"人民币", "美元", "港币"}
    currency_like_standard = currency in {"IFRS", "CAS", "PRC GAAP", "企业会计准则", "中国企业会计准则"}
    if standard and currency and standard_like_currency and currency_like_standard:
        add_note(
            notes,
            "逻辑与引用检查",
            "Medium",
            ws.title,
            "C7:C8",
            "适用会计准则与记账本位币疑似填反。",
            "会计准则和本位币字段语义错误会影响底稿基础信息、单位展示和后续交付一致性。",
            "SOP Lead Sheet 易错点：客户名称、期末、货币/单位等基础信息需更新，并与项目基础信息一致。",
            "请将 C7 填为适用会计准则（如 IFRS/CAS），将 C8 填为记账本位币或单位（如 CNY/人民币/万元）。",
        )


def check_summary(wb: Workbook, notes: list[ReviewNote]) -> None:
    if "汇总" not in wb.sheetnames:
        return
    ws = wb["汇总"]
    sap_yes_row = None
    tod_yes_row = None
    for row in range(3, ws.max_row + 1):
        program = norm(ws[f"C{row}"].value)
        page = norm(ws[f"F{row}"].value)
        if not program and not page:
            continue
        if not page:
            continue
        decision = ws[f"G{row}"].value
        reason = ws[f"H{row}"].value
        if page in {"程序页", "底稿索引", "N/A"}:
            continue
        if is_yes(decision) and "TOD" not in page.upper() and ("SAP" in program.upper() or "SAP" in page.upper() or "264GL" in program or "264GL" in page or "分析性复核" in program):
            sap_yes_row = row
        if is_yes(decision) and "TOD" in page.upper():
            tod_yes_row = row
        is_placeholder = any(token in page for token in ("项目组添加底稿索引", "如执行", "请添加", "请项目组"))
        if is_blank(decision):
            add_note(
                notes,
                "基础完整性检查",
                "Medium",
                ws.title,
                f"G{row}",
                "汇总页程序执行状态未选择。",
                "程序执行状态缺失会导致 Manager 无法判断该 PSP 是否已执行或已合理拒绝。",
                "SOP 汇总基础操作：确定执行的程序在 G 列选择“是”；拒绝执行在 G 列选择“否”并说明原因。",
                "请在 G 列选择“是/否”；如不执行，请在 H 列补充具体、可审阅的原因。",
            )
        if is_yes(decision) and is_placeholder:
            add_note(
                notes,
                "重点审计程序检查",
                "High",
                ws.title,
                f"F{row}",
                "汇总页标记执行程序，但底稿索引仍为模板占位说明。",
                "执行状态为“是”但未提供实际底稿索引，无法证明该审计程序已执行并形成可追踪证据链。",
                "SOP 汇总基础操作：确定执行的程序需在汇总页列明对应程序页或底稿索引。",
                "请将 F 列替换为实际 SAP/264GL/相关底稿索引；如实际不执行，请改为“否”并补充不执行原因。",
            )
        if is_no(decision) and is_blank(reason):
            add_note(
                notes,
                "重点审计程序检查",
                "High",
                ws.title,
                f"H{row}",
                "汇总页拒绝执行 PSP，但未填写不执行原因。",
                "拒绝 PSP 缺少原因可能意味着必要程序被遗漏，影响审计证据充分性。",
                "SOP 汇总易错点：拒绝 PSP 未解释原因，或解释不具体。",
                "请结合风险评估、账户性质和已执行替代程序，补充不执行原因及相关底稿索引。",
            )
        if is_yes(decision) and page and page in wb.sheetnames:
            target_ws = wb[page]
            if count_non_empty(target_ws) <= 8:
                add_note(
                    notes,
                    "重点审计程序检查",
                    "High",
                    ws.title,
                    f"F{row}",
                    f"汇总页标记执行“{page}”，但对应页签有效内容较少。",
                    "程序被标记为执行但底稿没有实质记录，可能无法支持审计结论。",
                    "SOP 汇总易错点：实际执行的程序与程序要求不一致，或少执行必要审计程序。",
                    "请补充该程序页的执行过程、证据索引和结论；如实际未执行，应更新汇总页并说明原因。",
                )
    if sap_yes_row and tod_yes_row:
        add_note(
            notes,
            "专业提示与优化建议",
            "Low",
            ws.title,
            f"G{sap_yes_row}:G{tod_yes_row}",
            "SAP 与 TOD 均标记执行，建议补充程序组合逻辑说明。",
            "SAP 与 TOD 可以同时执行，但若缺少组合逻辑说明，Manager 可能无法判断 SAP 证据是否不足以单独支持结论、为何仍需执行 TOD。",
            "SOP 汇总进阶提示：SAP 可减少 TOD 工作量；若 SAP 证据不足以就账户余额得出结论，则执行详细测试 TOD。",
            "请补充说明 SAP 是否已执行、其证据是否充分，以及继续执行 TOD 的原因；同时确保 SAP 和 TOD 均有清晰底稿索引。",
        )


def count_non_empty(ws) -> int:
    total = 0
    for row in ws.iter_rows():
        for cell in row:
            if not is_blank(cell.value):
                total += 1
    return total


def check_bkd(wb: Workbook, notes: list[ReviewNote], wb_values: Workbook | None = None) -> None:
    for sheet in ("VC.00 销售费用BKD", "VD.00 管理费用BKD"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        ws_values = wb_values[sheet] if wb_values is not None and sheet in wb_values.sheetnames else ws
        for row in range(8, 12):
            if is_blank(ws[f"C{row}"].value):
                add_note(
                    notes,
                    "重点审计程序检查",
                    "High",
                    sheet,
                    f"C{row}",
                    "CRA 等级未填写，测试阈值缺少风险评估基础。",
                    "CRA 缺失会影响测试阈值和后续程序性质、时间、范围的判断。",
                    "SOP BKD 易错点：各认定 CRA 正确，TT 取值正确；CRA 选择错误或未更新。",
                    "请根据 Canvas CRA 底稿补充对应认定的 CRA 等级，并复核 TT 计算是否合理。",
                )
        expectation_rows = find_rows_containing(ws, ("预期", "expectation"))
        if expectation_rows:
            start = expectation_rows[0] + 1
            sample_rows = list(range(start, min(start + 4, ws.max_row + 1)))
            if all(row_is_effectively_blank(ws, r, start_col=2, end_col=min(ws.max_column, 15)) for r in sample_rows):
                add_note(
                    notes,
                    "重点审计程序检查",
                    "Medium",
                    sheet,
                    f"B{sample_rows[0]}",
                    "BKD 预期或波动分析区域未见有效记录。",
                    "未建立预期或未记录理由，会削弱对异常波动识别和后续程序设计的支持。",
                    "SOP BKD 易错点：未记录波动预期；仅记录预期结果，未记录依据和理由。",
                    "请补充本期与上期变动预期、依据、异常波动判断，以及拟执行的跟进程序或索引。",
                )
        note_cells = find_cells_containing(ws, ("业务变化", "无异常", "正常波动", "无需进一步"))
        for coord, text in note_cells[:4]:
            if len(text) < 20:
                add_note(
                    notes,
                    "专业提示与优化建议",
                    "Low",
                    sheet,
                    coord,
                    "波动分析描述较笼统，缺少量化原因或程序索引。",
                    "过于概括的描述难以支持 Manager 判断异常波动是否已充分调查。",
                    "SOP ARP 易错点：仅描述变动金额与比例，未详细分析变动原因；仅分析、未说明程序性质及相关索引。",
                    "建议补充变动金额、比例、业务原因、检查的支持性文件，以及相关底稿索引。",
                )

        for row in range(1, min(ws.max_row, 35) + 1):
            for col in range(1, min(ws.max_column, 8) + 1):
                text = norm(ws_values.cell(row=row, column=col).value)
                if "货币" in text and "IFRS" in text.upper():
                    add_note(
                        notes,
                        "基础完整性检查",
                        "Medium",
                        sheet,
                        ws_values.cell(row=row, column=col).coordinate,
                        "BKD 表头货币/单位显示为 IFRS，疑似引用了会计准则字段。",
                        "货币/单位错误会影响底稿基础信息展示，并可能反映 Lead 基础字段引用错误。",
                        "SOP Lead Sheet 易错点：客户名称、期末、货币/单位等基础信息需准确更新。",
                        "请修正 Lead 中会计准则和记账本位币，并刷新 BKD 表头，使货币/单位显示为 CNY/人民币等正确单位。",
                    )
                    break

        threshold, pct_threshold = find_bkd_thresholds(ws_values)
        for row in range(25, min(ws.max_row, 120) + 1):
            amount_change = parse_number(ws_values.cell(row=row, column=12).value)
            pct_change = parse_number(ws_values.cell(row=row, column=13).value)
            quantitative_flag = ws_values.cell(row=row, column=16).value
            qualitative_flag = ws_values.cell(row=row, column=17).value
            account_name = norm(ws_values.cell(row=row, column=4).value)
            note_marker = norm(ws_values.cell(row=row, column=15).value)
            if (
                threshold is not None
                and amount_change is not None
                and abs(amount_change) > abs(threshold)
                and (pct_threshold is None or pct_change is None or abs(pct_change) > abs(pct_threshold))
                and not is_yes(quantitative_flag)
                and not is_yes(qualitative_flag)
            ):
                add_note(
                    notes,
                    "重点审计程序检查",
                    "High",
                    sheet,
                    f"P{row}:Q{row}",
                    f"{account_name or 'BKD 明细项'}变动超过波动阈值，但未标记进一步调查。",
                    "超过门槛的异常波动未调查，可能导致重大或性质异常的费用变动未被识别和跟进。",
                    "SOP BKD 基础操作：变动绝对值金额大于波动阈值且变动率超过范围的明细项应标记调查；P 或 Q 列为“是”的项目需填写 note 索引并执行 ARP 分析。",
                    "请将该项目标记为需进一步调查，补充变动原因、支持性证据和相关底稿索引；如判断无需调查，请记录具体理由。",
                )
            elif (
                threshold is not None
                and amount_change is not None
                and abs(amount_change) > abs(threshold)
                and "[A]" in note_marker
                and not is_yes(quantitative_flag)
                and not is_yes(qualitative_flag)
            ):
                add_note(
                    notes,
                    "重点审计程序检查",
                    "High",
                    sheet,
                    f"P{row}:Q{row}",
                    f"{account_name or 'BKD 明细项'}变动金额超过波动阈值且已关联波动说明，但未标记进一步调查。",
                    "金额超过门槛且已纳入波动说明的项目未在 P/Q 列标记调查，可能导致异常波动的 ARP 分析记录不完整。",
                    "SOP BKD 基础操作：超过门槛、性质异常或与预期不符的变动需调查，并对原因执行的调查结果详细记录、分析合理。",
                    "请将该项目标记为需进一步调查，或补充说明为何已通过相关测算/底稿索引充分应对而无需单独调查。",
                )

        for row in range(1, ws_values.max_row + 1):
            row_values = [ws_values.cell(row=row, column=col).value for col in range(1, min(ws_values.max_column, 18) + 1)]
            if any(norm(value).lower() == "diff" for value in row_values):
                numeric_values = [abs(float(value)) for value in row_values if isinstance(value, (int, float))]
                diff_amount = max(numeric_values) if numeric_values else 0
                if diff_amount:
                    add_note(
                        notes,
                        "逻辑与引用检查",
                        "High",
                        sheet,
                        f"A{row}:F{row}",
                        f"BKD 存在未解释的非零 Diff：{diff_amount:,.2f}。",
                        "BKD 明细与总账/Lead 勾稽存在差异且未说明，可能影响费用明细完整性和后续波动分析基础。",
                        "SOP Review Checklist：BKD 本期账面数与 TB/A3 核对一致，BKD 本期和上期审定数与 Lead 核对一致。",
                        "请补充差异构成说明、调整明细或扩大列示范围，确保 BKD 与总账/Lead 勾稽一致。",
                    )


def find_rows_containing(ws, keywords: Iterable[str]) -> list[int]:
    rows: list[int] = []
    lowers = tuple(k.lower() for k in keywords)
    for row in ws.iter_rows():
        for cell in row:
            text = norm(cell.value).lower()
            if text and any(k in text for k in lowers):
                rows.append(cell.row)
                break
    return rows


def find_cells_containing(ws, keywords: Iterable[str]) -> list[tuple[str, str]]:
    cells: list[tuple[str, str]] = []
    lowers = tuple(k.lower() for k in keywords)
    for row in ws.iter_rows():
        for cell in row:
            text = norm(cell.value)
            if text and any(k in text.lower() for k in lowers):
                cells.append((cell.coordinate, text))
    return cells


def row_is_effectively_blank(ws, row: int, start_col: int, end_col: int) -> bool:
    return all(is_blank(ws.cell(row=row, column=col).value) for col in range(start_col, end_col + 1))


def find_bkd_thresholds(ws) -> tuple[float | None, float | None]:
    amount_threshold = None
    pct_threshold = None
    for row in range(1, min(ws.max_row, 35) + 1):
        label = norm(ws.cell(row=row, column=2).value)
        value = parse_number(ws.cell(row=row, column=3).value)
        if "波动幅度" in label and "％" not in label and "%" not in label and "（%）" not in label and "(%)" not in label:
            amount_threshold = value
        if "波动幅度" in label and ("%" in label or "％" in label):
            pct_threshold = value
    return amount_threshold, pct_threshold


def check_tod(wb: Workbook, notes: list[ReviewNote], wb_values: Workbook | None = None) -> None:
    for sheet in ("VC&VD.01.2 TOD 年审", "VC&VD.01.2 TOD 预审+剩余期间"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        ws_values = wb_values[sheet] if wb_values is not None and sheet in wb_values.sheetnames else ws
        if sheet.endswith("年审"):
            required = [("D12", "样本总体名称"), ("E12", "数据来源 X-ref"), ("F12", "总体金额")]
            key_item_cells = [("D35", "关键项记录"), ("F35", "关键项金额")]
            sample_cells = [("D72", "样本检查记录"), ("H72", "支持性证据/结论")]
        else:
            required = [("D11", "样本总体名称"), ("E11", "数据来源 X-ref"), ("F11", "预审期间金额"), ("G11", "剩余期间金额")]
            key_item_cells = [("D35", "关键项记录"), ("F35", "关键项金额")]
            sample_cells = [("D83", "样本检查记录"), ("H83", "支持性证据/结论")]
        for coord, label in required:
            if "预审" in sheet and coord == "E11":
                continue
            if is_blank(ws[coord].value):
                add_note(
                    notes,
                    "重点审计程序检查",
                    "High",
                    sheet,
                    coord,
                    f"TOD {label} 未填写。",
                    "样本总体或来源缺失会影响抽样基础的完整性和可追溯性。",
                    "SOP TOD 基础操作：填列样本总体、数据来源、金额，并核对总体完整性。",
                    f"请补充 {label}，并与序时账/费用明细账或抽样工具输出建立交叉索引。",
                )
        if "预审" in sheet:
            missing_source_rows = [
                row for row in range(11, 15)
                if not is_blank(ws_values[f"D{row}"].value) and is_blank(ws_values[f"E{row}"].value)
            ]
            if len(missing_source_rows) > 1:
                add_note(
                    notes,
                    "重点审计程序检查",
                    "High",
                    sheet,
                    f"E{missing_source_rows[0]}:E{missing_source_rows[-1]}",
                    "TOD 多个样本总体的数据来源 X-ref 未填写。",
                    "样本总体缺少数据来源索引，会影响总体完整性、金额核对和抽样基础的可追溯性。",
                    "SOP TOD 基础操作：填列样本总体、数据来源、金额，并核对总体完整性。",
                    "请在各费用类别行补充对应序时账、费用明细账或样本池页签索引，并说明总体金额如何核对。",
                )
        for coord, label in key_item_cells + sample_cells:
            if coord in ws and is_blank(ws[coord].value):
                add_note(
                    notes,
                    "基础完整性检查",
                    "Medium",
                    sheet,
                    coord,
                    f"TOD {label} 未见记录。",
                    "关键项和样本检查记录不足，会影响细节测试是否覆盖高风险项目的判断。",
                    "SOP TOD 易错点：关键项、抽样策略、样本检查需记录依据和执行结果。",
                    f"请补充 {label}，包括选样依据、支持性证据、测试属性和结论。",
                )
        if "预审" in sheet:
            neg_amount = ws_values["D93"].value
            neg_sampling = ws_values["E93"].value
            if isinstance(neg_amount, (int, float)) and neg_amount < 0 and is_blank(neg_sampling):
                add_note(
                    notes,
                    "重点审计程序检查",
                    "Medium",
                    sheet,
                    "E93",
                    "负值样本总体存在金额，但是否抽样进行进一步工作的结构化字段未填写。",
                    "负值冲销或其他性质负值缺少明确字段结论或交叉索引，会影响费用完整性和异常冲销风险应对的可复核性。",
                    "SOP TOD 易错点：需分析负值冲销合理性，并记录负值样本抽样策略。",
                    "请在 E93 明确填写是否抽样；如相关判断已记录在其他位置，请补充“详见...”索引，并说明金额、SAD/TE 比较和替代程序。",
                )
            strategy = norm(ws_values["D49"].value)
            sample_amount = parse_number(ws_values["J78"].value)
            if ("1-9" in strategy or "1－9" in strategy) and ("1-10" in norm(ws_values["F10"].value) or "1-10" in norm(ws_values["F11"].value)):
                add_note(
                    notes,
                    "重点审计程序检查",
                    "Medium",
                    sheet,
                    "D49",
                    "TOD 抽样策略描述的预审/剩余期间口径与总体表头不一致。",
                    "期间拆分口径不一致可能导致样本总体、抽样工具输出和实际测试样本无法对应。",
                    "SOP TOD 抽样策略：应记录样本总体、抽样工具输出、样本选取结果，并与执行测试期间保持一致。",
                    "请统一预审和剩余期间口径；如实际为 1-10 月与 11-12 月，应修订 D49 文本并复核 Skywind 输出参数。",
                )
            elif "剩余期间" in strategy and "1" in strategy and sample_amount is not None and abs(sample_amount) < 1000:
                add_note(
                    notes,
                    "专业提示与优化建议",
                    "Low",
                    sheet,
                    "D49",
                    "剩余期间仅抽取 1 个代表性样本且样本金额较低，建议确认抽样工具输出是否已充分支持。",
                    "剩余期间样本数量和金额覆盖较低时，若缺少清晰说明，可能影响 TOD 代表性样本的证据充分性判断。",
                    "SOP TOD 抽样策略：应记录抽样工具输出、样本选取结果和不执行/限制样本量的理由。",
                    "请确认 Skywind 输出与样本分配一致；如样本量由工具计算得出，请保留并索引对应输出。",
                )

            population_names = [norm(ws_values[f"D{row}"].value) for row in range(11, 15)]
            non_vcvd = [name for name in population_names if any(token in name for token in ("研发", "制造"))]
            if non_vcvd:
                add_note(
                    notes,
                    "重点审计程序检查",
                    "Medium",
                    sheet,
                    "D11:H14",
                    "TOD 总体包含研发支出、制造费用等非 VC&VD 费用类别，但底稿范围未清晰说明。",
                    "不同交易类别未清晰区分样本总体，可能导致 VC&VD 程序范围、抽样总体和执行结论不一致。",
                    "SOP TOD Review Checklist：不同交易类别应区分不同样本总体；不建议将其他非本程序测试范围的交易明细一同上传后再剔除。",
                    "请明确本 TOD 是否覆盖全部 U_exp；如仅覆盖 VC&VD，请剔除研发/制造费用或建立单独底稿，并补充范围说明及索引。",
                )

            major_exclusions: list[str] = []
            for row in range(20, 28):
                name = norm(ws_values[f"D{row}"].value)
                xref = norm(ws_values[f"F{row}"].value)
                total = parse_number(ws_values[f"I{row}"].value)
                if total is not None and abs(total) > 0 and any(token in name for token in ("推广", "销售支持")):
                    if not xref or xref in {"　", "N/A", "NA"}:
                        major_exclusions.append(f"{name}（第{row}行）缺少执行程序索引")
                    else:
                        major_exclusions.append(f"{name}（第{row}行）为重大剔除项，需确认索引是否充分")
            if major_exclusions:
                add_note(
                    notes,
                    "重点审计程序检查",
                    "Medium",
                    sheet,
                    "D25:I26",
                    "推广费/销售支持服务费等重大剔除项需补充充分说明或审计程序索引。",
                    "重大剔除项若未清晰说明执行了何种程序，可能导致 TOD 样本总体被大额剔除但审计证据链不完整。",
                    "SOP TOD 易错点：剔除项应索引并 link 数据来源，重大剔除项应执行或索引恰当的审计程序，并记录关键项/剔除理由。",
                    "请确认并补充重大剔除项的数据来源、执行审计工作的具体底稿和剔除理由；若已执行测算或单独测试，请建立明确 X-ref。",
                )

    check_skywind_metadata(wb, notes, wb_values)
    if "VC&VD.01.2 详细测试 TOD" in wb.sheetnames:
        ws = wb["VC&VD.01.2 详细测试 TOD"]
        linked_text = " ".join(norm(c.value) for row in ws.iter_rows() for c in row if c.value)
        has_tod_workbook = any("TOD 年审" in s or "TOD 预审" in s for s in wb.sheetnames)
        if "TOD" in linked_text and not has_tod_workbook:
            add_note(
                notes,
                "重点审计程序检查",
                "Medium",
                ws.title,
                "B3",
                "主底稿仅链接 TOD 标准底稿，但当前文件中未包含实际 TOD 执行页。",
                "如果 TOD 被视为已执行但未上传或未纳入 Review 范围，Manager 无法复核样本测试过程。",
                "SOP 汇总易错点：索引至其他底稿的审计程序记录正确，相关底稿均已上传 Canvas。",
                "请上传 TOD 底稿一起 Review，或在主底稿中明确 TOD 文件索引和执行状态。",
            )


def check_skywind_metadata(wb: Workbook, notes: list[ReviewNote], wb_values: Workbook | None = None) -> None:
    value_book = wb_values or wb
    skywind_sheets = [ws for ws in value_book.worksheets if "Skywind" in ws.title and "Setting" not in ws.title]
    if not skywind_sheets:
        return

    customers: list[tuple[str, str, str]] = []
    tes: list[tuple[str, str, float]] = []
    scopes: list[tuple[str, str]] = []
    for ws in skywind_sheets:
        customer = norm(ws["D2"].value)
        te = parse_number(ws["F15"].value)
        scope = norm(ws["F12"].value)
        if customer:
            customers.append((ws.title, "D2", customer))
        if te is not None:
            tes.append((ws.title, "F15", te))
        if scope:
            scopes.append((ws.title, scope))

    unique_te = {round(item[2], 2) for item in tes}
    if len(unique_te) > 1:
        detail = "；".join(f"{sheet}!{cell}={value:,.2f}" for sheet, cell, value in tes)
        add_note(
            notes,
            "重点审计程序检查",
            "High",
            tes[-1][0],
            tes[-1][1],
            f"TOD 抽样工具中的 TE 不一致：{detail}。",
            "TE 参数不一致会直接影响关键项识别、样本量计算和抽样结论。",
            "SOP Lead/TOD 易错点：PM/TE/SAD 应与 Canvas 最终结果一致；抽样工具中的参数设置应正确。",
            "请统一预审、剩余期间及 Lead 使用的 TE，并重新评估 Skywind 抽样结果是否需要追加样本。",
        )

    expected_customer = ""
    if "Uexp.00 Lead" in value_book.sheetnames:
        expected_customer = norm(value_book["Uexp.00 Lead"]["C2"].value)
    elif customers:
        expected_customer = customers[0][2]
    for sheet, cell, customer in customers:
        if expected_customer and customer != expected_customer:
            add_note(
                notes,
                "基础完整性检查",
                "High",
                sheet,
                cell,
                f"Skywind 抽样底稿客户名称为“{customer}”，与当前底稿客户“{expected_customer}”不一致。",
                "客户名称错误可能表示模板复制未更新，影响底稿专业性和抽样输出归属判断。",
                "SOP Lead Sheet 易错点：客户名称、期末、货币/单位等基础信息需准确更新。",
                "请将 Skywind 抽样底稿客户名称更正为当前被审计单位，并确认其他基础信息同步更新。",
            )


def check_legal_fee(wb: Workbook, notes: list[ReviewNote]) -> None:
    sheet = "VD.01.3 复核法律费用"
    if sheet not in wb.sheetnames:
        return
    ws = wb[sheet]
    for row in range(12, 19):
        amount = ws[f"H{row}"].value or ws[f"E{row}"].value
        service = ws[f"F{row}"].value
        circular = ws[f"I{row}"].value
        if not is_blank(amount) and is_blank(service):
            add_note(
                notes,
                "基础完整性检查",
                "Medium",
                sheet,
                f"F{row}",
                "法律费用存在金额，但服务性质/内容未填写。",
                "缺少服务性质会影响对诉讼、罚款、律师费等特殊费用的风险识别。",
                "SOP ARP 易错点：诉讼费/罚款/律师费等性质特殊的费用，需分析业务背景。",
                "请补充服务内容、业务背景和支持文件索引。",
            )
        if not is_blank(amount) and is_blank(circular):
            add_note(
                notes,
                "重点审计程序检查",
                "Medium",
                sheet,
                f"I{row}",
                "法律费用存在金额，但是否发函或不发函原因未记录。",
                "法律费用可能涉及未决诉讼或或有事项，未说明发函判断会影响风险应对。",
                "SOP 法律费用程序：对性质特殊的法律费用向律师发送审计询问函；不发函需列示原因。",
                "请记录是否发函；如不发函，请说明替代程序和判断依据。",
            )


def check_cutoff(wb: Workbook, notes: list[ReviewNote], wb_values: Workbook | None = None) -> None:
    sheet = "VC&VD.01.4 截止性测试"
    if sheet not in wb.sheetnames:
        return
    ws = wb[sheet]
    ws_values = wb_values[sheet] if wb_values is not None and sheet in wb_values.sheetnames else ws
    for coord, label, severity in [
        ("B10", "测试期间天数", "High"),
        ("C10", "测试期间确定依据/原因", "High"),
        ("D10", "测试开始日期", "Medium"),
        ("E10", "测试结束日期", "Medium"),
    ]:
        if is_blank(ws[coord].value):
            add_note(
                notes,
                "重点审计程序检查",
                severity,
                sheet,
                coord,
                f"截止性测试“{label}”未填写。",
                "测试期间或依据缺失会影响期末前后费用是否记录在正确期间的判断。",
                "SOP 截止性测试：结合费用发生到入账周期确定测试期间，并记录确定依据。",
                f"请补充 {label}，并说明测试期间如何覆盖期末前后风险。",
            )
    strategy = norm(ws_values["B13"].value)
    lead_te = None
    if wb_values is not None and "Uexp.00 Lead" in wb_values.sheetnames:
        lead_te = parse_number(wb_values["Uexp.00 Lead"]["C5"].value)
    sample_amounts: list[float] = []
    for row in range(24, min(ws.max_row, 45) + 1):
        amount_value = parse_number(ws_values[f"D{row}"].value)
        if amount_value is not None and not is_blank(ws_values[f"B{row}"].value):
            sample_amounts.append(abs(amount_value))
    if lead_te is not None and "大于TE" in strategy.replace(" ", "").upper() and sample_amounts and max(sample_amounts) < abs(lead_te):
        add_note(
            notes,
            "重点审计程序检查",
            "High",
            sheet,
            "B13",
            "截止性测试策略描述为筛选发生额大于 TE 的交易，但实际样本金额均低于 TE。",
            "选样策略与实际执行不一致，会导致 Manager 无法判断截止性测试是否按既定标准覆盖关键样本。",
            "SOP 截止性测试：应采用判断抽样方式选取关键项和代表性样本，并在底稿中恰当记录选取过程和结果。",
            "请修正策略描述为实际执行口径，或按策略重新选取大于 TE 的样本；如采用前 N 大/代表性样本，请补充原因。",
        )

    for row in range(24, min(ws.max_row, 45) + 1):
        voucher = ws_values[f"B{row}"].value
        amount = ws_values[f"D{row}"].value
        evidence = ws_values[f"F{row}"].value
        ey_works = ws_values[f"K{row}"].value
        remark = ws_values[f"L{row}"].value
        is_sample_row = isinstance(amount, (int, float)) and not is_blank(voucher)
        if is_sample_row and is_blank(evidence):
            add_note(
                notes,
                "重点审计程序检查",
                "Medium",
                sheet,
                f"F{row}",
                "截止性测试样本存在金额，但支持性证据信息未填写。",
                "缺少支持性证据会影响样本是否已记录在适当会计期间的判断。",
                "SOP 截止性测试：根据支持性证据确定所选交易已记录在适当会计期间。",
                "请补充发票、结算单、合同、付款单据等支持性证据信息和索引。",
            )
        if is_sample_row and is_blank(ey_works):
            add_note(
                notes,
                "基础完整性检查",
                "Medium",
                sheet,
                f"K{row}",
                "截止性测试样本未记录 EY Works 属性测试结果。",
                "缺少属性测试结果会影响样本是否已记录在适当会计期间的判断。",
                "SOP 截止性测试要求记录执行过程和结果。",
                "请在 K 列记录测试结果；如存在例外或为 N，请在 L 列补充异常解释和处理结论。",
            )
        if is_sample_row and norm(ey_works).upper() in {"N", "NO", "否"} and is_blank(remark):
            add_note(
                notes,
                "重点审计程序检查",
                "Medium",
                sheet,
                f"L{row}",
                "截止性测试样本存在异常结果，但备注未解释原因和处理。",
                "异常结果缺少解释会影响截止性差异是否已充分调查和解决。",
                "SOP 截止性测试要求对例外情况记录如何解决并执行额外程序。",
                "请在 L 列说明异常原因、影响金额、是否调整及额外审计程序。",
            )


def check_professional_wording(wb: Workbook, notes: list[ReviewNote]) -> None:
    highlight_fill_tokens = {"FFFF00", "FFFFFF00", "00FFFF00", "FFFF99", "FFFFFF99"}
    for ws in wb.worksheets:
        if not is_audit_sheet(ws.title):
            continue
        for row in ws.iter_rows():
            for cell in row:
                fill = cell.fill
                if fill and fill.fill_type == "solid":
                    color = fill.fgColor.rgb or fill.fgColor.indexed
                    if str(color).upper() in highlight_fill_tokens:
                        add_note(
                            notes,
                            "专业提示与优化建议",
                            "Low",
                            ws.title,
                            cell.coordinate,
                            "底稿中存在黄色高亮单元格，可能为未清理的待办或 Review 痕迹。",
                            "交付前未清理 highlight 可能表示底稿仍存在未解决事项。",
                            "SOP 汇总易错点：底稿中不存在重大 OS 事项，review notes 均已解决并删除，底稿中无 highlight。",
                            "请确认该高亮是否仍需保留；如事项已解决，请清理高亮并更新结论。",
                        )


def export_notes(notes: list[ReviewNote], source_name: str) -> tuple[Path, Path]:
    ensure_dirs()
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_path = OUTPUT_DIR / f"review_notes_{Path(source_name).stem}_{stamp}.xlsx"
    csv_path = OUTPUT_DIR / f"review_notes_{Path(source_name).stem}_{stamp}.csv"

    from openpyxl import Workbook as NewWorkbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = NewWorkbook()
    ws = wb.active
    ws.title = "Review Notes"
    headers = ["编号", "分类", "严重程度", "来源", "置信度", "Sheet", "Cell/区域", "问题描述", "审计风险", "依据", "建议修改"]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for note in notes:
        ws.append([
            note.id,
            note.category,
            note.severity,
            note.source,
            note.confidence,
            note.sheet,
            note.cell,
            note.issue,
            note.audit_risk,
            note.basis,
            note.suggestion,
        ])
    widths = [10, 22, 12, 12, 10, 28, 14, 38, 42, 42, 42]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sev = row[2].value
        fill = {"High": "F8CBAD", "Medium": "FFE699", "Low": "D9EAD3"}.get(sev)
        if fill:
            row[2].fill = PatternFill("solid", fgColor=fill)
    ws.freeze_panes = "A2"
    wb.save(xlsx_path)

    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for note in notes:
            writer.writerow([note.id, note.category, note.severity, note.source, note.confidence, note.sheet, note.cell, note.issue, note.audit_risk, note.basis, note.suggestion])
    return xlsx_path, csv_path


def highlight_workbook(source: Path, notes: list[ReviewNote]) -> Path:
    ensure_dirs()
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = OUTPUT_DIR / f"highlighted_{source.stem}_{stamp}.xlsx"
    shutil.copy2(source, output)
    wb = load_workbook(output)
    fill_map = {
        "High": PatternFill("solid", fgColor="F4B084"),
        "Medium": PatternFill("solid", fgColor="FFD966"),
        "Low": PatternFill("solid", fgColor="C6E0B4"),
    }
    for note in notes:
        if note.sheet not in wb.sheetnames:
            continue
        if not note.cell or ":" in note.cell:
            continue
        try:
            cell = wb[note.sheet][note.cell]
        except Exception:
            continue
        from openpyxl.cell.cell import Cell
        if not isinstance(cell, Cell):
            continue
        cell.fill = copy(fill_map.get(note.severity, fill_map["Medium"]))
        cell.comment = Comment(f"{note.id} {note.issue}\n建议：{note.suggestion}", "AI Review")
    add_navigation_sheet(wb, notes)
    wb.save(output)
    return output


def add_navigation_sheet(wb: Workbook, notes: list[ReviewNote]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    sheet_name = "底稿导航"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name, 0)

    title_fill = PatternFill("solid", fgColor="1F4E79")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    header_fill = PatternFill("solid", fgColor="EAF2F8")
    border = Border(bottom=Side(style="thin", color="D9E2F3"))

    ws["A1"] = "底稿导航"
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=16)
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A1:H1")
    ws.row_dimensions[1].height = 28

    ws["A3"] = "一、底稿总览"
    ws["A3"].font = Font(bold=True, color="1F4E79")
    ws["A3"].fill = section_fill
    ws.merge_cells("A3:H3")

    overview_headers = ["序号", "底稿模块", "Sheet名称", "底稿用途", "关键定位", "点击跳转", "主要内容提示", "来源"]
    ws.append(overview_headers)
    style_header_row(ws, 4, header_fill, border)

    row_idx = 5
    visible_sheets = [s for s in wb.sheetnames if s != sheet_name and is_audit_sheet(s)]
    for index, target_sheet in enumerate(visible_sheets, 1):
        key_cell = "A1"
        ws.append([
            index,
            classify_sheet_module(target_sheet),
            target_sheet,
            describe_sheet_purpose(target_sheet),
            key_cell,
            "打开",
            build_sheet_hint(target_sheet),
            "原始底稿",
        ])
        add_internal_link(ws.cell(row=row_idx, column=6), target_sheet, key_cell)
        row_idx += 1

    issue_start = row_idx + 2
    ws.cell(row=issue_start, column=1).value = "二、异常直达"
    ws.cell(row=issue_start, column=1).font = Font(bold=True, color="1F4E79")
    ws.cell(row=issue_start, column=1).fill = section_fill
    ws.merge_cells(start_row=issue_start, start_column=1, end_row=issue_start, end_column=8)

    issue_headers = ["风险等级", "问题编号", "分类", "所属Sheet", "问题摘要", "定位单元格", "点击跳转", "建议处理"]
    header_row = issue_start + 1
    for col, value in enumerate(issue_headers, 1):
        ws.cell(row=header_row, column=col).value = value
    style_header_row(ws, header_row, header_fill, border)

    row_idx = header_row + 1
    for note in notes:
        ws.append([
            note.severity,
            note.id,
            note.category,
            note.sheet,
            note.issue,
            note.cell or "Sheet",
            "定位",
            note.suggestion,
        ])
        if note.sheet in wb.sheetnames:
            add_internal_link(ws.cell(row=row_idx, column=7), note.sheet, note.cell or "A1")
        severity_fill = {"High": "F8CBAD", "Medium": "FFE699", "Low": "D9EAD3"}.get(note.severity)
        if severity_fill:
            ws.cell(row=row_idx, column=1).fill = PatternFill("solid", fgColor=severity_fill)
        row_idx += 1

    widths = [10, 20, 30, 38, 18, 14, 34, 46]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    ws.freeze_panes = "A5"


def style_header_row(ws, row: int, fill: PatternFill, border: Border) -> None:
    from openpyxl.styles import Alignment, Font

    for cell in ws[row]:
        cell.font = Font(bold=True, color="1F1F1F")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border


def add_internal_link(cell, sheet_name: str, target_cell: str) -> None:
    safe_sheet = sheet_name.replace("'", "''")
    safe_cell = target_cell if target_cell and ":" not in target_cell else "A1"
    cell.hyperlink = f"#'{safe_sheet}'!{safe_cell}"
    cell.style = "Hyperlink"


def classify_sheet_module(sheet_name: str) -> str:
    if "Lead" in sheet_name:
        return "基础信息"
    if "BKD" in sheet_name:
        return "明细核对"
    if "TOD" in sheet_name:
        return "细节测试"
    if "SAP" in sheet_name.upper() or "汇总" in sheet_name:
        return "程序汇总"
    if "法律" in sheet_name or "娉曞緥" in sheet_name:
        return "专项复核"
    if "截止" in sheet_name or "鎴" in sheet_name:
        return "截止性测试"
    return "审计底稿"


def describe_sheet_purpose(sheet_name: str) -> str:
    if "Lead" in sheet_name:
        return "记录项目基础信息、期间、准则、币种及阈值等关键参数。"
    if "BKD" in sheet_name:
        return "核对费用明细、波动、差异及相关解释。"
    if "TOD" in sheet_name:
        return "记录细节测试样本、抽样范围、执行结果和剩余期间程序。"
    if "法律" in sheet_name or "娉曞緥" in sheet_name:
        return "复核法律费用及相关支持性证据。"
    if "截止" in sheet_name or "鎴" in sheet_name:
        return "检查期前期后费用确认是否记录在恰当期间。"
    if "汇总" in sheet_name:
        return "汇总审计程序执行情况、索引和结论。"
    return "审计底稿工作页。"


def build_sheet_hint(sheet_name: str) -> str:
    if "Lead" in sheet_name:
        return "重点看公司名称、期间、准则、币种、TE/PM/SAD。"
    if "BKD" in sheet_name:
        return "重点看 Diff、波动解释、重大金额和异常标记。"
    if "TOD" in sheet_name:
        return "重点看样本覆盖、抽样判断、负值样本和执行结论。"
    if "截止" in sheet_name or "鎴" in sheet_name:
        return "重点看样本日期、金额阈值、例外事项和备注解释。"
    return "点击进入该底稿页查看。"


def create_demo_workbook() -> Path:
    ensure_dirs()
    target = SAMPLE_DIR / "模拟问题底稿_AIReviewDemo.xlsx"
    templates = sorted(BASE_DIR.glob("模板-U_exp SWP VC&VD*.xlsx"))
    if not templates:
        raise FileNotFoundError("未找到标准费用底稿模板。")
    shutil.copy2(templates[0], target)
    wb = load_workbook(target)

    lead = wb["Uexp.00 Lead"]
    lead["C2"] = "XYZ公司"
    lead["C3"] = datetime(2025, 12, 31)
    lead["C4"] = None
    lead["C5"] = None
    lead["C6"] = None
    lead["C7"] = "CAS"
    lead["C8"] = None

    summary = wb["汇总"]
    summary["G3"] = "是"
    summary["G5"] = "是"
    summary["G8"] = "否"
    summary["H8"] = None
    summary["G10"] = "是"
    summary["G12"] = "否"
    summary["H12"] = None
    summary["G14"] = "是"

    for sheet in ("VC.00 销售费用BKD", "VD.00 管理费用BKD"):
        ws = wb[sheet]
        for row in range(8, 12):
            ws[f"C{row}"] = None
        vague_cell = first_writable_cell(ws, 18, 30, 15, 22)
        vague_cell.value = "业务变化"
        vague_cell.fill = PatternFill("solid", fgColor="FFFF00")
        ws["D8"] = "=#REF!"

    legal = wb["VD.01.3 复核法律费用"]
    legal["D12"] = "律师咨询费"
    legal["E12"] = 120000
    legal["F12"] = None
    legal["H12"] = 120000
    legal["I12"] = None

    cutoff = wb["VC&VD.01.4 截止性测试"]
    cutoff["B10"] = None
    cutoff["C10"] = None
    cutoff["I24"] = 88000
    cutoff["F24"] = None

    wb.save(target)
    return target


def first_writable_cell(ws, start_row: int, end_row: int, start_col: int, end_col: int):
    from openpyxl.cell.cell import Cell

    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell, Cell):
                return cell
    raise ValueError(f"未找到可写入单元格：{ws.title}")


def as_public_payload(notes: list[ReviewNote], stats: dict, source_path: Path) -> dict:
    xlsx_path, csv_path = export_notes(notes, source_path.name)
    highlighted_path = highlight_workbook(source_path, notes)
    return {
        "stats": stats,
        "notes": [asdict(n) | {"location": n.location} for n in notes],
        "exports": {
            "xlsx": str(xlsx_path),
            "csv": str(csv_path),
            "highlighted": str(highlighted_path),
        },
        "source": scan_workbook(source_path),
        "checkpoints": load_checkpoints(),
    }
