from __future__ import annotations

import json
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from llm_client import LLMUnavailable, chat_json
from review_engine import OUTPUT_DIR, ReviewNote, build_stats, load_checkpoints, scan_workbook


CATEGORIES = {"基础完整性检查", "重点审计程序检查", "逻辑与引用检查", "专业提示与优化建议"}
SEVERITIES = {"High", "Medium", "Low"}
ACTIONS = {"keep", "revise", "drop", "add"}


SYSTEM_PROMPT = """你是 FSO Audit CoE 的 Manager，负责基于 VC&VD SOP 复核费用类审计底稿。
你的任务是降低规则脚本误判、补充有依据的漏判，并把 Review Notes 改写成专业、可执行的审阅意见。

硬性要求：
1. 只能依据输入中的 workbook 摘要、规则命中、SOP checkpoint 和 sheet/cell 证据判断，不得编造。
2. 没有 sheet/cell 或明确区域依据的问题，不得新增为 Review Note。
3. 不要把“程序页”“执行”“不执行的原因”等表头当成缺失页签。
4. TOD 单独底稿不应被要求包含主 SWP 的全部页签。
5. 外部底稿链接、说明文字、模板占位符不直接判为程序缺失；只有汇总页明确标记执行且当前 workbook 应存在对应页签时才提示。
6. 输出必须是 JSON object，且只包含一个 key: notes。
7. notes 必须是数组，只返回需要变更的项：drop、revise、add。没有问题的规则项必须省略，系统会自动保留。
8. 每个对象字段固定为：id, category, severity, sheet, cell, issue, audit_risk, basis, suggestion, confidence, action。
9. action 只能是 revise、drop、add；除非确需覆盖置信度，否则不要输出 keep。
10. revise/drop 必须引用原始规则 id；add 可使用空 id。
11. category 只能是：基础完整性检查、重点审计程序检查、逻辑与引用检查、专业提示与优化建议。
12. severity 只能是 High、Medium、Low。
13. 严禁把原始规则结果原样复制成 revise；只有分类、严重程度、定位、风险判断或建议有实质改动时才 revise。
14. 如未发现误判或漏判，返回 {"notes": []}。
15. 发现误判必须 action=drop；发现重复事项必须保留一条、drop 其余，不得用 revise 规避 drop。
16. notes 最多返回 25 条，排序必须是 drop 在前、add 其次、revise 最后。
17. 本产品采用 recall-first 口径：规则命中的基础信息错误、PSP 索引缺失、波动阈值未调查、TE/客户名称不一致、负值判断空白、BKD Diff、截止性策略与执行不一致等硬事实，除非输入证据明确反驳，不得 drop。
18. 允许保留 benchmark 之外的合理 findings；不要为了减少数量而删除有 SOP 依据、sheet/cell 清晰的问题。
19. 对 TOD 总体范围、重大剔除项说明这类语义判断，可在有明确单元格证据时 add 或 revise；语气应为“需补充说明/索引”，避免过度断言底稿一定错误。
"""


def enhance_review_with_llm(source_path: Path, rule_notes: list[ReviewNote], rule_stats: dict) -> tuple[list[ReviewNote], dict]:
    meta = {
        "mode": "rule",
        "message": "LLM 未启用，当前为规则结果。",
        "debug_path": "",
    }
    try:
        payload = build_llm_payload(source_path, rule_notes, rule_stats)
        response = chat_json(
            [
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
            ],
            temperature=0.0,
            max_tokens=8192,
        )
        final_notes = apply_llm_actions(rule_notes, response)
        debug_path = write_debug_payload(source_path, payload, response, final_notes)
        meta = {
            "mode": "llm",
            "message": "已完成 DeepSeek LLM 增强 Review。",
            "debug_path": str(debug_path),
        }
        return final_notes, meta
    except LLMUnavailable as exc:
        meta["message"] = str(exc)
        return rule_notes, meta
    except Exception as exc:
        meta["message"] = f"LLM 输出校验失败，已回退到规则结果：{exc}"
        return rule_notes, meta


def build_llm_payload(source_path: Path, rule_notes: list[ReviewNote], rule_stats: dict) -> dict:
    return {
        "task": "请复核规则引擎输出，删除误判、修正分类/严重程度/措辞，并在有明确证据时补充漏判。",
        "workbook": workbook_summary(source_path),
        "scan": scan_workbook(source_path),
        "rule_stats": rule_stats,
        "rule_notes": [compact_note(note) for note in rule_notes],
        "checkpoints": load_checkpoints(),
        "output_contract": {
            "object_key": "notes",
            "fields": ["id", "category", "severity", "sheet", "cell", "issue", "audit_risk", "basis", "suggestion", "confidence", "action"],
        },
    }


def compact_note(note: ReviewNote) -> dict:
    return {
        "id": note.id,
        "category": note.category,
        "severity": note.severity,
        "sheet": note.sheet,
        "cell": note.cell,
        "issue": note.issue,
        "basis": note.basis,
        "location": note.location,
    }


def workbook_summary(path: Path, max_cells_per_sheet: int = 45) -> dict:
    wb = load_workbook(path, data_only=False, read_only=True)
    sheets: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        if ws.title.startswith(("Skywind", "DS_INTERNAL")):
            continue
        cells = []
        for row in ws.iter_rows():
            for cell in row:
                append_cell(cells, cell)
                if len(cells) >= max_cells_per_sheet:
                    break
            if len(cells) >= max_cells_per_sheet:
                break
        key_cells = targeted_cells(ws)
        sheets.append({"name": ws.title, "rows": ws.max_row, "cols": ws.max_column, "sample_cells": cells, "key_cells": key_cells})
    return {"file": path.name, "sheets": sheets}


def append_cell(cells: list[dict], cell) -> None:
    if cell.value is None:
        return
    text = str(cell.value).replace("\n", " ").strip()
    if not text:
        return
    cells.append({"cell": cell.coordinate, "value": text[:180]})


def targeted_cells(ws) -> list[dict]:
    title = ws.title
    ranges: list[tuple[int, int, int, int]] = []
    if title == "汇总":
        ranges.append((2, min(ws.max_row, 18), 2, 8))
    elif title == "Uexp.00 Lead":
        ranges.append((2, 8, 2, 3))
    elif "BKD" in title:
        ranges.extend([(17, 24, 2, 17), (55, min(ws.max_row, 75), 2, 17)])
    elif "截止性测试" in title:
        ranges.extend([(8, 15, 2, 5), (23, min(ws.max_row, 45), 2, 12)])
    elif "法律费用" in title:
        ranges.append((12, min(ws.max_row, 22), 2, 12))
    elif "TOD" in title:
        ranges.extend([(9, 35, 3, 10), (45, 52, 3, 10), (76, 82, 3, 10), (90, 99, 3, 10)])

    cells: list[dict] = []
    seen: set[str] = set()
    for min_row, max_row, min_col, max_col in ranges:
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                before = len(cells)
                append_cell(cells, cell)
                if len(cells) > before:
                    if cells[-1]["cell"] in seen:
                        cells.pop()
                    else:
                        seen.add(cells[-1]["cell"])
    return cells[:160]


def apply_llm_actions(rule_notes: list[ReviewNote], response: Any) -> list[ReviewNote]:
    items = response.get("notes") if isinstance(response, dict) else response
    if not isinstance(items, list):
        raise ValueError("LLM JSON 中缺少 notes 数组。")
    by_id = {note.id: note for note in rule_notes}
    decisions: dict[str, dict] = {}
    additions: list[ReviewNote] = []

    for raw in items:
        if not isinstance(raw, dict):
            continue
        action = str(raw.get("action", "")).strip()
        if action not in ACTIONS:
            continue
        original_id = str(raw.get("id", "")).strip()
        if action in {"drop", "keep", "revise"} and original_id in by_id:
            decisions[original_id] = raw
        elif action == "add":
            note = note_from_llm(raw)
            if note.sheet and note.cell:
                additions.append(note)

    kept: list[ReviewNote] = []
    for note in rule_notes:
        decision = decisions.get(note.id)
        if not decision:
            kept.append(note)
            continue
        action = str(decision.get("action", "")).strip()
        if action == "drop":
            continue
        kept.append(note_from_llm(decision, fallback=note))
    kept.extend(additions)
    return renumber_notes(dedupe_notes(kept))


def note_from_llm(raw: dict, fallback: ReviewNote | None = None) -> ReviewNote:
    def pick(field: str, default: str = "") -> str:
        value = raw.get(field)
        if value is None and fallback is not None:
            value = getattr(fallback, field)
        return str(value or default).strip()

    category = pick("category", fallback.category if fallback else "专业提示与优化建议")
    severity = pick("severity", fallback.severity if fallback else "Low")
    if category not in CATEGORIES:
        category = fallback.category if fallback else "专业提示与优化建议"
    if severity not in SEVERITIES:
        severity = fallback.severity if fallback else "Low"
    confidence = parse_confidence(raw.get("confidence", getattr(fallback, "confidence", 0.75) if fallback else 0.75))
    source = "llm" if fallback is None else "rule+llm"
    return ReviewNote(
        id=pick("id", fallback.id if fallback else ""),
        category=category,
        severity=severity,
        sheet=pick("sheet"),
        cell=pick("cell"),
        issue=pick("issue"),
        audit_risk=pick("audit_risk"),
        basis=pick("basis"),
        suggestion=pick("suggestion"),
        confidence=confidence,
        source=source,
    )


def parse_confidence(value) -> float:
    if isinstance(value, (int, float)):
        return max(0.0, min(1.0, float(value)))
    text = str(value or "").strip().lower()
    mapping = {"high": 0.9, "medium": 0.75, "low": 0.6, "高": 0.9, "中": 0.75, "低": 0.6}
    if text in mapping:
        return mapping[text]
    try:
        return max(0.0, min(1.0, float(text)))
    except ValueError:
        return 0.75


def dedupe_notes(notes: list[ReviewNote]) -> list[ReviewNote]:
    result: list[ReviewNote] = []
    seen: set[tuple[str, ...]] = set()
    for note in notes:
        issue_text = f"{note.sheet} {note.cell} {note.issue} {note.suggestion}"
        if "TOD" in issue_text and any(token in issue_text for token in ("链接", "外部", "执行页", "执行记录", "标准底稿", "有效内容较少", "实质记录")):
            key = ("TOD_EXTERNAL_WORKPAPER",)
        else:
            key = (note.sheet, note.cell, note.category, note.issue[:30])
        if key in seen:
            continue
        seen.add(key)
        result.append(note)
    return result


def renumber_notes(notes: list[ReviewNote]) -> list[ReviewNote]:
    for idx, note in enumerate(notes, 1):
        note.id = f"RN-{idx:03d}"
    return notes


def write_debug_payload(source_path: Path, request_payload: dict, response: Any, final_notes: list[ReviewNote]) -> Path:
    debug_dir = OUTPUT_DIR / "llm_debug"
    debug_dir.mkdir(parents=True, exist_ok=True)
    path = debug_dir / f"llm_debug_{source_path.stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    safe_payload = {
        "request": request_payload,
        "response": response,
        "final_notes": [asdict(n) for n in final_notes],
    }
    path.write_text(json.dumps(safe_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def final_stats(notes: list[ReviewNote]) -> dict:
    return build_stats(notes)
