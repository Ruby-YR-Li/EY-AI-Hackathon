import unittest
import tempfile
from pathlib import Path
from unittest.mock import patch

import pandas as pd

from src.ai_client import AIJudgeResult
from src.ai_prompts import AI_CHECK_TARGETS
from src.checker import QualityChecker
from src.models import AuditWorkbook, CellLocation, CheckCategory, CheckResult, RuleStatus, Severity
from src.rules.base import BaseRule
from src.rules.completeness import CompletenessRules
from src.rules.consistency import ConsistencyRules
from src.rules.format import FormatRules
from src.rules.logic import LogicRules


class RuntimeErrorRule(BaseRule):
    def __init__(self):
        super().__init__("TEST-001", "测试规则")

    def list_rule_ids(self):
        return [{
            "rule_id": "TEST-001",
            "rule_name": "测试规则",
            "category": CheckCategory.BASIC_PROC,
            "required_sheets": [],
        }]

    def check(self, workbook):
        self.reset_runtime_state()
        self.mark_rule_error("TEST-001", "模拟子规则异常")
        return []


class RuntimeSkipOnlyRule(BaseRule):
    def __init__(self):
        super().__init__("AI-SKIP", "AI跳过规则")

    def list_rule_ids(self):
        return [{
            "rule_id": "AI-SKIP",
            "rule_name": "AI跳过规则",
            "category": CheckCategory.BASIC_PROC,
            "required_sheets": [],
        }]

    def check(self, workbook):
        self.reset_runtime_state()
        self.mark_rule_skipped("AI-SKIP", "AI调用失败，未完成AI判断: 模拟")
        return []


class RuntimeSkipWithIssueRule(BaseRule):
    def __init__(self):
        super().__init__("AI-PARTIAL", "混合AI规则")

    def list_rule_ids(self):
        return [{
            "rule_id": "AI-PARTIAL",
            "rule_name": "混合AI规则",
            "category": CheckCategory.BASIC_PROC,
            "required_sheets": [],
        }]

    def check(self, workbook):
        self.reset_runtime_state()
        self.mark_rule_skipped("AI-PARTIAL", "AI调用失败，未完成AI判断: 模拟")
        return [CheckResult(
            rule_id="AI-PARTIAL",
            rule_name="混合AI规则",
            category=CheckCategory.BASIC_PROC,
            severity=Severity.HIGH,
            location=CellLocation("Sheet1", 1, 1),
            message="Python确定性问题",
            expected="无问题",
            actual="存在问题",
        )]


class ErrorAIClient:
    def judge(self, prompt):
        return AIJudgeResult(True, "模拟AI连接失败", "人工复核", is_error=True)


class RuleExecutionSafetyTests(unittest.TestCase):
    def test_registered_rule_ids_match_checklist_checkpoints(self):
        checklist_path = Path("资料库") / "检查规则清单.xlsx"
        checklist = pd.read_excel(checklist_path, header=0)
        checklist_ids = {
            str(rule_id).strip()
            for rule_id in checklist["规则ID"]
            if str(rule_id).strip() and str(rule_id).strip().lower() != "nan"
        }

        registered_ids = set()
        for rule in [
            FormatRules(enable_ai=False),
            ConsistencyRules(),
            LogicRules(enable_ai=False),
            CompletenessRules(enable_ai=False),
        ]:
            registered_ids.update(rdef["rule_id"] for rdef in rule.list_rule_ids())
        registered_ids.update(AI_CHECK_TARGETS)

        self.assertEqual(checklist_ids, registered_ids)

    def test_sop_format_rules_are_registered_and_executed(self):
        rule = FormatRules(enable_ai=False)

        registered_ids = {rdef["rule_id"] for rdef in rule.list_rule_ids()}

        self.assertIn("SOP-FMT001", registered_ids)
        self.assertIn("SOP-FMT002", registered_ids)

    def test_sop_fmt001_reports_missing_basic_info_value(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["客户名称", ""],
                    ["适用会计准则", "企业会计准则"],
                    ["记账本位币", "人民币"],
                ]),
            },
        )
        rule = FormatRules(enable_ai=False)

        results = rule._check_basic_info(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SOP-FMT001", results[0].rule_id)
        self.assertIn("客户名称", results[0].message)

    def test_sop_fmt002_requires_yyyy_mm_dd_string_dates(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["期末日期", "2026/03/31"],
                    ["分析日期", "2026-03-31"],
                ]),
            },
        )
        rule = FormatRules(enable_ai=False)

        results = rule._check_date_format(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SOP-FMT002", results[0].rule_id)
        self.assertIn("期末日期格式", results[0].message)

    def test_sub_rule_exception_is_not_reported_as_passed(self):
        checker = QualityChecker(enable_ai=False)
        checker.rules = [RuntimeErrorRule()]
        checker.parser.parse = lambda _: AuditWorkbook(file_path="test.xlsx", sheets={})
        checker.parser.validate_structure = lambda _: []

        _, _, records = checker.check("test.xlsx")
        record = next(r for r in records if r.rule_id == "TEST-001")

        self.assertEqual(RuleStatus.FAILED, record.status)
        self.assertIn("模拟子规则异常", record.skipped_reason)

    def test_passed_rule_keeps_actual_parser_evidence(self):
        checker = QualityChecker(enable_ai=False)
        checker.ai_rules = None

        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["名义金额", 100],
                ]),
                "K.01 Agree SL to GL": pd.DataFrame([
                    ["TB-原值", None, None, None, None],
                    [None, None, None, None, 0],
                ]),
            },
        )
        checker.parser.parse = lambda _: workbook

        _, _, records = checker.check("test.xlsx")
        gl001 = next(record for record in records if record.rule_id == "GL-001")

        self.assertEqual(RuleStatus.PASSED, gl001.status)
        evidence = {
            (loc.sheet, loc.row, loc.col)
            for loc in gl001.evidence_locations
        }
        self.assertIn(("K.00 Lead Sheet", 1, 2), evidence)
        self.assertIn(("K.01 Agree SL to GL", 2, 5), evidence)

    def test_unrecognized_key_structures_are_reported_as_review_issues(self):
        cases = [
            (
                "summary_without_psp_columns",
                {"汇总": pd.DataFrame([["无关", "内容"]])},
                ["AE-003", "DP-003"],
            ),
            (
                "k01_without_gl_areas",
                {
                    "K.00 Lead Sheet": pd.DataFrame([["名义金额", 100]]),
                    "K.01 Agree SL to GL": pd.DataFrame([["无关", 0]]),
                },
                ["GL-001", "GL-002", "GL-003b", "AE-004c"],
            ),
            (
                "k00_without_investigation_header",
                {"K.00 Lead Sheet": pd.DataFrame([["无关", 0]])},
                ["AE-004c"],
            ),
            (
                "policy_sheet_without_policy_range",
                {
                    "K.03.3 折旧政策复核": pd.DataFrame([["无关", 0]]),
                    "FA list": pd.DataFrame([
                        ["资产类别", "资产名称", "使用年限"],
                        ["机器", "机器A", 5],
                    ]),
                },
                ["DP-002"],
            ),
        ]

        for case_name, sheets, rule_ids in cases:
            with self.subTest(case=case_name):
                checker = QualityChecker(enable_ai=False)
                checker.ai_rules = None
                workbook = AuditWorkbook(file_path="case.xlsx", sheets=sheets)
                checker.parser.parse = lambda _: workbook
                checker.parser.validate_structure = lambda _: []

                _, _, records = checker.check("case.xlsx")
                by_rule = {record.rule_id: record for record in records}
                for rule_id in rule_ids:
                    self.assertEqual(
                        RuleStatus.FAILED,
                        by_rule[rule_id].status,
                        f"{case_name} {rule_id} should be 发现问题",
                    )
                    self.assertEqual(1, len(by_rule[rule_id].issues))
                    self.assertEqual(Severity.REVIEW, by_rule[rule_id].issues[0].severity)
                    self.assertIn("需人工复核", by_rule[rule_id].issues[0].message)

    def test_ae006_manual_review_prompt_uses_review_severity(self):
        rule = ConsistencyRules()
        workbook = AuditWorkbook(file_path="test.xlsx", sheets={})

        results = rule._check_cross_references(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("AE-006", results[0].rule_id)
        self.assertEqual(Severity.REVIEW, results[0].severity)

    def test_ae001_reports_external_materiality_baseline_mismatch(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["整体重要性", 1000],
                    ["可容忍误差", 800],
                    ["名义金额", 100],
                ]),
            },
            context={
                "baseline": {
                    "pm": {"value": 1100},
                    "te": {"value": 900},
                    "sad": {"value": 120},
                }
            },
        )
        rule = FormatRules(enable_ai=False)

        results = rule._check_te_sad(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("AE-001", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)
        self.assertIn("PM(1000.0)与项目配置PM(1100)不一致", results[0].message)
        self.assertIn("TE(800.0)与基准TE(900)不一致", results[0].message)
        self.assertIn("SAD(100.0)与基准SAD(120)不一致", results[0].message)

    def test_gl003_reports_lead_a3_baseline_mismatch(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["项目", "期末账面数", "期末审定数", "上期末审定数"],
                    ["原值", 100, 90, 80],
                ]),
            },
            context={
                "baseline": {
                    "a3": {
                        "cost": {
                            "book_amount": {"value": 100},
                            "audited_amount": {"value": 91},
                            "prior_audited_amount": {"value": 80},
                        }
                    }
                }
            },
        )
        rule = ConsistencyRules()

        results = rule._check_lead_vs_a3(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("GL-003", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)
        self.assertIn("K.00 原值期末审定数(90.0)与A3(91.0)不一致", results[0].message)

    def test_gl001_bkd_tb_difference_over_sad_is_reported(self):
        k01_rows = [[""] * 18 for _ in range(4)]
        k01_rows[0][0] = "TB-原值"
        k01_rows[1][4] = 200
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([["名义金额", 100]]),
                "K.01 Agree SL to GL": pd.DataFrame(k01_rows),
            },
        )
        rule = ConsistencyRules()

        results = rule._check_bkd_vs_tb(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("GL-001", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)

    def test_gl002_bkd_falist_difference_over_sad_is_reported(self):
        k01_rows = [[""] * 10 for _ in range(4)]
        k01_rows[0][0] = "表2 check with 表1"
        k01_rows[1][1] = 200
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([["名义金额", 100]]),
                "K.01 Agree SL to GL": pd.DataFrame(k01_rows),
            },
        )
        rule = ConsistencyRules()

        results = rule._check_bkd_vs_falist(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("GL-002", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)

    def test_gl003b_k01_audit_total_vs_k00_difference_is_reported(self):
        k01_rows = [[""] * 7 for _ in range(6)]
        k01_rows[0][4] = "审定数"
        k01_rows[0][5] = "审定数"
        k01_rows[1][0] = "原值"
        k01_rows[2][0] = "年初余额"
        k01_rows[2][4] = 80
        k01_rows[2][5] = 0
        k01_rows[3][0] = "年末余额"
        k01_rows[3][4] = 90
        k01_rows[3][5] = 0
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["项目", "期末审定数", "上期末审定数"],
                    ["原值", 91, 80],
                ]),
                "K.01 Agree SL to GL": pd.DataFrame(k01_rows),
            },
        )
        rule = ConsistencyRules()

        results = rule._check_k01_vs_k00_audit(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("GL-003b", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)

    def test_gl004_depreciation_vs_pnl_difference_uses_high_severity(self):
        k01_rows = [[""] * 10 for _ in range(8)]
        k01_rows[0][0] = "表4"
        k01_rows[2][4] = "差异"
        k01_rows[2][5] = 200
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([["名义金额", 100]]),
                "K.01 Agree SL to GL": pd.DataFrame(k01_rows),
            },
        )
        rule = ConsistencyRules()

        results = rule._check_depreciation_vs_pnl(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("GL-004", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)

    def test_gl004_finds_diff_label_outside_fixed_column_and_right_shifted_amount(self):
        k01_rows = [[""] * 14 for _ in range(8)]
        k01_rows[0][0] = "表4 折旧费用与利润表核对"
        k01_rows[3][2] = "区域差异项"
        k01_rows[3][11] = 200
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([["名义金额", 100]]),
                "K.01 Agree SL to GL": pd.DataFrame(k01_rows),
            },
        )
        rule = ConsistencyRules()

        results = rule._check_depreciation_vs_pnl(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("GL-004", results[0].rule_id)
        self.assertEqual(12, results[0].location.col)

    def test_gl004_parses_formatted_amount_text(self):
        k01_rows = [[""] * 10 for _ in range(8)]
        k01_rows[0][0] = "表4"
        k01_rows[2][4] = "差异"
        k01_rows[2][5] = "￥（1，234.56）"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([["名义金额", 100]]),
                "K.01 Agree SL to GL": pd.DataFrame(k01_rows),
            },
        )
        rule = ConsistencyRules()

        results = rule._check_depreciation_vs_pnl(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("GL-004", results[0].rule_id)
        self.assertEqual("￥（1，234.56）", results[0].actual)

    def test_gl005_abnormal_movement_uses_high_severity(self):
        k01_rows = [[""] * 8 for _ in range(5)]
        k01_rows[0][0] = "原值"
        k01_rows[1][0] = "购置"
        k01_rows[1][4] = -1
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.01 Agree SL to GL": pd.DataFrame(k01_rows)},
        )
        rule = ConsistencyRules()

        results = rule._check_abnormal_movements(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("GL-005", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)

    def test_gl005d_reports_k01_rollforward_mismatch(self):
        k01_rows = [[""] * 8 for _ in range(16)]
        k01_rows[0][4] = "账面数"
        k01_rows[2][2] = "年初余额"
        k01_rows[2][4] = 100
        k01_rows[3][2] = "购置"
        k01_rows[3][4] = 20
        k01_rows[4][1] = "原值"
        k01_rows[4][2] = "处置或报废"
        k01_rows[4][4] = 0
        k01_rows[5][2] = "年末余额"
        k01_rows[5][4] = 119
        k01_rows[6][2] = "年初余额"
        k01_rows[6][4] = 10
        k01_rows[7][2] = "计提"
        k01_rows[7][4] = 5
        k01_rows[8][1] = "累计折旧"
        k01_rows[8][2] = "处置或报废"
        k01_rows[8][4] = 0
        k01_rows[9][2] = "年末余额"
        k01_rows[9][4] = 15
        k01_rows[10][2] = "年初余额"
        k01_rows[10][4] = 1
        k01_rows[11][2] = "计提"
        k01_rows[11][4] = 0
        k01_rows[12][1] = "减值准备"
        k01_rows[12][2] = "处置或报废"
        k01_rows[12][4] = 0
        k01_rows[13][2] = "年末余额"
        k01_rows[13][4] = 1
        k01_rows[14][1] = "净值 (NBV)"
        k01_rows[14][2] = "年初余额"
        k01_rows[14][4] = 89
        k01_rows[15][2] = "年末余额"
        k01_rows[15][4] = 103
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.01 Agree SL to GL": pd.DataFrame(k01_rows)},
        )
        rule = ConsistencyRules()

        results = rule._check_k01_rollforward_completeness(workbook)

        self.assertTrue(any(result.rule_id == "GL-005d" for result in results))
        self.assertTrue(any("原值" in result.rule_name for result in results))
        self.assertTrue(any("重算120.00" in result.actual for result in results))

    def test_ae005_unreadable_workbook_is_marked_skipped(self):
        workbook = AuditWorkbook(file_path="missing.xlsx", sheets={})
        rule = FormatRules(enable_ai=False)
        rule.reset_runtime_state()

        with patch("src.rules.format.load_workbook", side_effect=OSError("cannot open")):
            results = rule._check_workbook_cleanliness(workbook)

        self.assertEqual([], results)
        self.assertIn("AE-005", rule.runtime_skips)
        self.assertIn("无法读取Excel批注信息", rule.runtime_skips["AE-005"])

    def test_runtime_skip_without_issue_is_reported_as_review_issue(self):
        checker = QualityChecker(enable_ai=False)
        checker.rules = [RuntimeSkipOnlyRule()]
        checker.parser.parse = lambda _: AuditWorkbook(file_path="test.xlsx", sheets={})
        checker.parser.validate_structure = lambda _: []

        _, _, records = checker.check("test.xlsx")
        record = next(r for r in records if r.rule_id == "AI-SKIP")

        self.assertEqual(RuleStatus.FAILED, record.status)
        self.assertIn("AI调用失败，未完成AI判断", record.skipped_reason)
        self.assertEqual(1, len(record.issues))
        self.assertEqual(Severity.REVIEW, record.issues[0].severity)

    def test_runtime_skip_with_python_issue_keeps_failed_issue_and_reason(self):
        checker = QualityChecker(enable_ai=False)
        checker.rules = [RuntimeSkipWithIssueRule()]
        checker.parser.parse = lambda _: AuditWorkbook(
            file_path="test.xlsx",
            sheets={"Sheet1": pd.DataFrame([["x"]])},
        )
        checker.parser.validate_structure = lambda _: []

        _, results, records = checker.check("test.xlsx")
        record = next(r for r in records if r.rule_id == "AI-PARTIAL")

        self.assertEqual(1, len(results))
        self.assertEqual(RuleStatus.FAILED, record.status)
        self.assertIn("AI调用失败，未完成AI判断", record.skipped_reason)
        self.assertEqual(1, len(record.issues))
        self.assertEqual("Python确定性问题", record.issues[0].message)

    def test_ae004b_ai_error_is_marked_skipped_not_business_issue(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["预期及额外考虑", "本期固定资产变动与业务计划一致"],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=True)
        rule.ai_client = ErrorAIClient()
        rule.reset_runtime_state()

        results = rule._check_expectation_reasonableness(workbook)

        self.assertEqual([], results)
        self.assertIn("AE-004b", rule.runtime_skips)
        self.assertIn("AI调用失败，未完成AI判断", rule.runtime_skips["AE-004b"])

    def test_ae004d_ai_error_is_marked_skipped_not_business_issue(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["波动说明", "本期固定资产增加主要来自正常生产投入"],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=True)
        rule.ai_client = ErrorAIClient()
        rule.reset_runtime_state()

        results = rule._check_fluctuation_explanation_quality(workbook)

        self.assertEqual([], results)
        self.assertIn("AE-004d", rule.runtime_skips)
        self.assertIn("AI调用失败，未完成AI判断", rule.runtime_skips["AE-004d"])

    def test_ae004c_missing_notes_marker_is_reported(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["项目", "notes", "变动金额", "变动%", "进一步调查"],
                    ["原值", "", 1000, "50%", "是"],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_fluctuation_notes_marker(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("AE-004c", results[0].rule_id)
        self.assertIn("notes列为空", results[0].message)

    def test_ae004_empty_fluctuation_explanation_is_reported(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["波动说明", ""],
                    ["", ""],
                    ["", ""],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_fluctuation_investigation(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("AE-004", results[0].rule_id)
        self.assertIn("波动说明", results[0].message)

    def test_sp003c_ai_error_is_marked_skipped_not_business_issue(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.1 新增测试": pd.DataFrame([
                    ["选择关键项目的理由", ""],
                    ["金额超过TT的项目全部作为关键项目测试", ""],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=True)
        rule.ai_client = ErrorAIClient()
        rule.reset_runtime_state()

        results = rule._check_addition_key_item_reason_quality(workbook)

        self.assertEqual([], results)
        self.assertIn("SP-003c", rule.runtime_skips)
        self.assertIn("AI调用失败，未完成AI判断", rule.runtime_skips["SP-003c"])

    def test_sp006_blank_disposal_key_item_reason_is_reported(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.2 处置测试": pd.DataFrame([
                    ["选择关键项目的理由", ""],
                    ["", ""],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_disposal_key_items(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-006", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)
        self.assertIn("选择关键项目的理由为空", results[0].message)

    def test_sp006a_ai_error_is_marked_skipped_not_business_issue(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.2 处置测试": pd.DataFrame([
                    ["选择关键项目的理由", ""],
                    ["处置金额超过TT的项目全部作为关键项目测试", ""],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=True)
        rule.ai_client = ErrorAIClient()
        rule.reset_runtime_state()

        results = rule._check_disposal_key_item_reason_quality(workbook)

        self.assertEqual([], results)
        self.assertIn("SP-006a", rule.runtime_skips)
        self.assertIn("AI调用失败，未完成AI判断", rule.runtime_skips["SP-006a"])

    def test_dt001_disposal_sample_mismatch_is_review_severity(self):
        headers = ["固定资产编号", "资产名称", "类别", "处置日期", "原值", "累计折旧", "减值准备", "净值"]
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.2": pd.DataFrame([
                    headers,
                    ["A001", "设备A", "机器设备", "2026-03-31", 100, 20, 0, 80],
                ]),
                "K.02.2a": pd.DataFrame([
                    headers,
                    ["A002", "设备B", "机器设备", "2026-03-31", 100, 20, 0, 80],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_disposal_sample_consistency(workbook)

        self.assertEqual(2, len(results))
        self.assertTrue(all(result.rule_id == "DT-001" for result in results))
        self.assertTrue(all(result.severity == Severity.REVIEW for result in results))

    def test_disposal_breakdown_issue_is_attributed_to_sp004(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.2": pd.DataFrame([
                    ["Breakdown", "", ""],
                    [100, "", ""],
                ]),
                "K.00 Lead Sheet": pd.DataFrame([
                    ["名义金额", 10],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_disposal_breakdown(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-004", results[0].rule_id)

    def test_unreadable_addition_samples_are_marked_skipped(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.1": pd.DataFrame([["无法识别的新增测试内容"]]),
                "K.02.1a": pd.DataFrame([["无法识别的选样内容"]]),
            },
        )
        rule = CompletenessRules(enable_ai=False)
        rule.reset_runtime_state()

        results = rule._check_addition_sample_consistency(workbook)

        self.assertEqual([], results)
        self.assertIn("AT-001", rule.runtime_skips)

    def test_addition_placeholder_none_is_not_treated_as_sample(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.1": pd.DataFrame([
                    ["固定资产编号", "资产名称", "类别", "原值"],
                    ["无", "无", "无", ""],
                ]),
                "K.02.1a": pd.DataFrame([
                    ["固定资产编号", "资产名称", "类别", "原值"],
                    ["A001", "生产设备", "机器设备", 100],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)
        rule.reset_runtime_state()

        results = rule._check_addition_sample_consistency(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("AT-001", results[0].rule_id)
        self.assertEqual(Severity.REVIEW, results[0].severity)
        self.assertEqual("选样输出存在，K.02.1缺失", results[0].actual)
        self.assertNotIn("'无'", results[0].message)

    def test_sampling_te_allows_rounding_tail_under_one_yuan(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([["可容忍误差", 304906.88]]),
                "K.02.1a": pd.DataFrame([["可容忍误差", 304907.25]]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_sampling_params_for_output(
            workbook, "K.02.1a", "SP-002", "新增"
        )

        self.assertEqual([], results)

    def test_sampling_te_flags_difference_above_one_yuan(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([["可容忍误差", 304906.88]]),
                "K.02.1a": pd.DataFrame([["可容忍误差", 304908.25]]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_sampling_params_for_output(
            workbook, "K.02.1a", "SP-002", "新增"
        )

        self.assertEqual(1, len(results))
        self.assertEqual("SP-002", results[0].rule_id)

    def test_sp001_addition_breakdown_diff_over_sad_is_reported(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["名义金额", 100],
                ]),
                "K.02.1": pd.DataFrame([
                    ["Breakdown", "购置"],
                    ["差异", 500],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_addition_breakdown(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-001", results[0].rule_id)
        self.assertIn("SAD=100", results[0].message)

    def test_sp001b_finds_sample_pool_amount_far_right(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.1": pd.DataFrame([
                    ["Breakdown中购置金额", "", "", "", "", 0],
                ]),
                "K.02.1a": pd.DataFrame([
                    ["样本池总体金额", "", "", "", "", 1000],
                    ["可容忍误差", "", "", "", "", 100],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_addition_sample_pool(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-001b", results[0].rule_id)
        self.assertIn("选样=1000", results[0].actual)

    def test_sp003_empty_addition_key_item_reason_is_reported(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.1 新增测试": pd.DataFrame([
                    ["选择关键项目的理由", ""],
                    ["", ""],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_addition_key_items(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-003", results[0].rule_id)
        self.assertIn("关键项选取理由为空", results[0].rule_name)

    def test_sp003a_finds_key_item_amount_on_same_row(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.1": pd.DataFrame([
                    ["减：测试的关键项目", "", "", "", "", 0],
                ]),
                "K.02.1a": pd.DataFrame([
                    ["定量关键项金额（金额大于100）", "", "", "", "", 1000],
                    ["可容忍误差", "", "", "", "", 100],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_addition_key_item_amount_match_sample(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-003a", results[0].rule_id)
        self.assertIn("K.02.1=0", results[0].actual)
        self.assertIn("选样输出=1000", results[0].actual)

    def test_sp003b_uses_addition_list_amounts_over_tt(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["存在", "", 500],
                ]),
                "K.02.1": pd.DataFrame([
                    ["减：测试的关键项目", "", "", "", "", 0],
                ]),
                "新增清单": pd.DataFrame([
                    ["资产名称", "原值-本期增加"],
                    ["A", 1000],
                    ["B", 400],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_addition_key_item_amount_match_tt(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-003b", results[0].rule_id)
        self.assertIn("超TT合计=1000", results[0].actual)

    def test_sp003b_blank_key_item_amount_can_represent_zero(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["存在", "", 500],
                ]),
                "K.02.1": pd.DataFrame([
                    ["减：测试的关键项目", "", "", "", "", ""],
                ]),
                "新增清单": pd.DataFrame([
                    ["资产名称", "原值-本期增加"],
                    ["A", 100],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)
        rule.reset_runtime_state()

        results = rule._check_addition_key_item_amount_match_tt(workbook)

        self.assertEqual([], results)
        self.assertNotIn("SP-003b", rule.runtime_skips)

    def test_sp004b_flags_disposal_sample_pool_mismatch(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.2": pd.DataFrame([
                    ["Breakdown", "", "", 800],
                ]),
                "K.02.2a": pd.DataFrame([
                    ["样本池总体金额", 1000],
                    ["可容忍误差", 100],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_disposal_sample_pool(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-004b", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)
        self.assertIn("选样=1000", results[0].actual)
        self.assertIn("K.02.2=800", results[0].actual)

    def test_sp005_flags_disposal_sampling_te_mismatch(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["可容忍误差", 500],
                ]),
                "K.02.2a": pd.DataFrame([
                    ["可容忍误差", 600],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_disposal_sampling_params(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-005", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)
        self.assertIn("处置选样输出", results[0].message)

    def test_sp006b_flags_disposal_representative_sample_count_mismatch(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.2": pd.DataFrame([
                    ["选择关键项目的理由", "计划测试2个代表性样本"],
                    ["样本类型", "资产名称"],
                    ["代表性样本", "设备A"],
                    ["代表性样本", "设备B"],
                ]),
                "K.02.2a": pd.DataFrame([
                    ["代表性样本量", 3],
                    ["样本类型", "资产名称"],
                    ["代表性样本", "设备A"],
                    ["代表性样本", "设备B"],
                    ["代表性样本", "设备C"],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_disposal_representative_sample_count(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-006b", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)
        self.assertIn("理由=2", results[0].actual)
        self.assertIn("抽样输出=3", results[0].actual)
        self.assertIn("测试sheet=2", results[0].actual)

    def test_mt002_reports_special_nature_movement_in_original_cost_section(self):
        rows = [[""] * 6 for _ in range(8)]
        rows[0][0] = "原值"
        rows[1][0] = "购置"
        rows[1][3] = 100
        rows[2][0] = "债务重组减少"
        rows[2][3] = -80
        rows[3][0] = "累计折旧"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.01 Agree SL to GL": pd.DataFrame(rows)},
        )
        rule = CompletenessRules(enable_ai=False)

        findings = rule._find_special_nature_changes(workbook)
        results = rule._check_special_nature_movements(findings)

        self.assertEqual(1, len(results))
        self.assertEqual("MT-002", results[0].rule_id)
        self.assertIn("债务重组减少", results[0].message)
        self.assertIn("减少", results[0].message)

    def test_at003_reports_special_nature_addition_in_original_cost_section(self):
        rows = [[""] * 6 for _ in range(8)]
        rows[0][0] = "原值"
        rows[1][0] = "投资者投入"
        rows[1][3] = 120
        rows[2][0] = "累计折旧"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.01 Agree SL to GL": pd.DataFrame(rows)},
        )
        rule = CompletenessRules(enable_ai=False)

        findings = rule._find_special_nature_changes(workbook)
        results = rule._check_addition_special_nature(findings)

        self.assertEqual(1, len(results))
        self.assertEqual("AT-003", results[0].rule_id)
        self.assertIn("投资者投入", results[0].message)

    def test_mt003_reports_missing_adjustment_detail(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["科目名称", "账表调整数", "审计调整数"],
                    ["原值", 0, 100],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_adjustment_items(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("MT-003", results[0].rule_id)
        self.assertEqual("调整明细缺失", results[0].rule_name)

    def test_dp003_reports_when_no_depreciation_test_strategy_executed(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "汇总": pd.DataFrame([
                    ["K.03程序", "执行", "不执行原因"],
                    ["K.03.1 SAP", "否", ""],
                    ["K.03.2 TOD", "否", ""],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_depreciation_strategy(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-003", results[0].rule_id)
        self.assertIn("均未标记执行", results[0].message)

    def test_dp004_ignores_investigation_label_when_locating_sap_difference(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.1 SAP": pd.DataFrame([
                    ["差异是否需要进一步调查", "否"],
                    ["折旧测试差异", 100],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_sap_depreciation(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-004", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)
        self.assertIn("差异=100", results[0].actual)

    def test_dp004a_flags_sap_te_mismatch_with_lead(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([["可容忍误差", 100]]),
                "K.03.1 SAP": pd.DataFrame([
                    ["实体类型", "非上市"],
                    ["CRA", "Minimal"],
                    ["TE", 200],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_sap_basic_parameters(workbook)

        self.assertTrue(any(result.rule_id == "DP-004a" for result in results))
        self.assertTrue(any("TE" in result.message for result in results))

    def test_dp004b_flags_missing_sap_expectation(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.03.1 SAP": pd.DataFrame([["实体类型", "非上市"], ["CRA", "Minimal"]])},
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_sap_expectation_sufficiency(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-004b", results[0].rule_id)

    def test_dp004c_flags_insufficient_precision_factors(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.03.1 SAP": pd.DataFrame([["预期精确度理由", "参考历史折旧趋势"]])},
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_sap_expectation_precision_factors(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-004c", results[0].rule_id)

    def test_dp004d_flags_threshold_over_te_limit(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([["可容忍误差", 100]]),
                "K.03.1 SAP": pd.DataFrame([
                    ["偏差阈值", 300],
                    ["分解说明", "按类别分解"],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_sap_threshold_and_decomposition(workbook)

        self.assertTrue(any(result.rule_id == "DP-004d" for result in results))
        self.assertTrue(any("超过2.5倍TE" in result.message for result in results))

    def test_dp004e_flags_sap_base_data_mismatch_with_bkd(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.01 Agree SL to GL": pd.DataFrame([
                    ["累计折旧", "账面数"],
                    ["本年计提", 100],
                ]),
                "K.03.1 SAP": pd.DataFrame([
                    ["本期折旧", 150],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_sap_base_data_vs_bkd(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-004e", results[0].rule_id)

    def test_dp004f_flags_unexplained_sap_difference(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.03.1 SAP": pd.DataFrame([["折旧测试差异", 100]])},
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_sap_difference_follow_up(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-004f", results[0].rule_id)

    def test_dp004g_flags_missing_sap_conclusion(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.03.1 SAP": pd.DataFrame([["测试结论", ""]])},
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_sap_conclusion_completeness(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-004g", results[0].rule_id)

    def test_dp005d_flags_special_formula_case_keywords(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.2 ByItem折旧测试": pd.DataFrame([
                    ["资产编号", "减值", "本期应折旧金额"],
                    ["A001", 10, 100],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_byitem_special_formula_cases(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-005d", results[0].rule_id)

    def test_dp005d_does_not_flag_impairment_header_when_values_are_zero(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.2 ByItem折旧测试": pd.DataFrame([
                    ["资产编号", "减值准备", "本期应折旧金额"],
                    ["A001", 0, 100],
                    ["A002", 0, 200],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_byitem_special_formula_cases(workbook)

        self.assertEqual([], results)

    def test_dp005_skips_intro_difference_text_and_uses_real_header(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["名义金额", 50],
                    ["资产负债表日", "2026-03-31"],
                ]),
                "K.03.2 ByItem折旧测试": pd.DataFrame([
                    ["折旧差异分析说明", "", "", ""],
                    ["", "", "", ""],
                    ["资产编号", "本期应折旧金额", "本期计提折旧", "本期差异"],
                    ["A001", "1,000", 900, 0],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_byitem_depreciation(workbook)

        self.assertTrue(any(result.rule_id == "DP-005" for result in results))
        self.assertTrue(any("重算不一致" in result.message for result in results))
        self.assertTrue(any("超过SAD" in result.rule_name for result in results))

    def test_dp005_uses_depreciation_amount_not_period_column(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["名义金额", 50],
                ]),
                "K.03.2 ByItem折旧测试": pd.DataFrame([
                    ["资产编号", "本期应折旧时间", "本期应折旧金额", "本期计提折旧", "本期差异"],
                    ["A001", 1, 100, 90, 10],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_byitem_depreciation(workbook)

        self.assertEqual([], results)

    def test_dp005a_compares_byitem_accrued_amount_with_k01_bkd(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["名义金额", 50],
                ]),
                "K.01 Agree SL to GL": pd.DataFrame([
                    ["累计折旧", "", "", ""],
                    ["本年计提", 50, 50, ""],
                ]),
                "K.03.2 ByItem折旧测试": pd.DataFrame([
                    ["资产编号", "本期应折旧金额", "本期计提折旧", "本期差异"],
                    ["A001", 100, 90, 10],
                    ["A002", 100, 90, 10],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_byitem_population_vs_bkd(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-005a", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)
        self.assertIn("By Item=180.00", results[0].actual)
        self.assertIn("K.01/BKD=100.00", results[0].actual)

    def test_dp005b_reports_dead_number_in_byitem_recalculation_column(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "K.03.2 ByItem折旧测试"
        headers = ["资产编号", "原值", "使用寿命(月)", "残值率", "本期计提折旧", "本期应折旧金额", "本期差异"]
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=header)
        ws.cell(row=2, column=1, value="A001")
        ws.cell(row=2, column=2, value=1000)
        ws.cell(row=2, column=3, value=10)
        ws.cell(row=2, column=4, value=0.05)
        ws.cell(row=2, column=5, value=90)
        ws.cell(row=2, column=6, value=95)
        ws.cell(row=2, column=7, value="=F2-E2")

        with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
            wb.save(tmp.name)
            workbook = AuditWorkbook(
                file_path=tmp.name,
                sheets={
                    "K.03.2 ByItem折旧测试": pd.DataFrame([
                        headers,
                        ["A001", 1000, 10, 0.05, 90, 95, 5],
                    ]),
                },
            )
            rule = LogicRules(enable_ai=False)

            results = rule._check_byitem_formula_policy(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-005b", results[0].rule_id)
        self.assertEqual(Severity.MEDIUM, results[0].severity)
        self.assertIn("非公式死数", results[0].message)

    def test_sp002c_reports_sampling_exclusion_without_reason(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.1a 新增选样输出": pd.DataFrame([
                    ["抽样参数", ""],
                    ["剔除项", ""],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_sampling_exclusions_for_output(
            workbook, "K.02.1a 新增选样输出", "SP-002c", "新增"
        )

        self.assertEqual(1, len(results))
        self.assertEqual("SP-002c", results[0].rule_id)
        self.assertEqual(Severity.REVIEW, results[0].severity)
        self.assertIn("缺少原因", results[0].rule_name)

    def test_sp005c_accepts_explicit_no_sampling_exclusion(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.2a 处置选样输出": pd.DataFrame([
                    ["抽样参数", ""],
                    ["无剔除项", ""],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_sampling_exclusions_for_output(
            workbook, "K.02.2a 处置选样输出", "SP-005c", "处置"
        )

        self.assertEqual([], results)
        self.assertGreater(rule.get_evidence_location_count("SP-005c"), 0)

    def test_sp002c_treats_zero_exclusion_total_as_no_exclusion(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.1a 新增选样输出": pd.DataFrame([
                    ["排除项总量", 0, 0, 0],
                    ["已删除负余额（行数）", 0, 0, 0],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_sampling_exclusions_for_output(
            workbook, "K.02.1a 新增选样输出", "SP-002c", "新增"
        )

        self.assertEqual([], results)
        self.assertGreater(rule.get_evidence_location_count("SP-002c"), 0)

    def test_at002a_collects_continued_evidence_lines(self):
        rows = [[""] * 18 for _ in range(6)]
        rows[0][14] = "EY Works"
        rows[0][15] = "EY Works"
        rows[1][12] = "获得的证据/支持的描述"
        rows[1][14] = "1"
        rows[1][15] = "2"
        rows[2][4] = "A001"
        rows[2][12] = "合同编号: XS-001"
        rows[2][14] = "Y"
        rows[2][15] = "Y"
        rows[3][12] = "验收单编号: YS-001"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.02.1 新增测试": pd.DataFrame(rows)},
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_addition_capitalization_evidence(workbook)

        self.assertEqual([], results)
        self.assertGreater(rule.get_evidence_location_count("AT-002a"), 0)

    def test_at002e_reports_generic_supporting_file_description(self):
        rows = [[""] * 18 for _ in range(4)]
        rows[0][14] = "EY Works"
        rows[1][12] = "获得的证据/支持的描述"
        rows[1][14] = "1"
        rows[2][4] = "A001"
        rows[2][12] = "采购单、发票"
        rows[2][14] = "Y"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.02.1 新增测试": pd.DataFrame(rows)},
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_addition_supporting_file_key_info(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("AT-002e", results[0].rule_id)
        self.assertEqual(Severity.REVIEW, results[0].severity)

    def test_at002c_reports_missing_depreciation_policy_evidence(self):
        rows = [[""] * 20 for _ in range(4)]
        rows[0][14] = "EY Works"
        rows[1][12] = "获得的证据/支持的描述"
        rows[1][14] = "1"
        rows[1][15] = "2"
        rows[1][16] = "3"
        rows[1][17] = "4"
        rows[2][4] = "A001"
        rows[2][12] = "合同、发票、验收单"
        rows[2][14] = "Y"
        rows[2][15] = "Y"
        rows[2][16] = "Y"
        rows[2][17] = "Y"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.02.1 新增测试": pd.DataFrame(rows)},
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_addition_depreciation_policy_evidence(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("AT-002c", results[0].rule_id)
        self.assertIn("属性3", results[0].message)

    def test_at002d_reports_missing_asset_category_evidence(self):
        rows = [[""] * 20 for _ in range(4)]
        rows[0][14] = "EY Works"
        rows[1][12] = "获得的证据/支持的描述"
        rows[1][14] = "1"
        rows[1][15] = "2"
        rows[1][16] = "3"
        rows[1][17] = "4"
        rows[2][4] = "A001"
        rows[2][12] = "使用寿命与折旧政策一致"
        rows[2][14] = "Y"
        rows[2][15] = "Y"
        rows[2][16] = "Y"
        rows[2][17] = "Y"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.02.1 新增测试": pd.DataFrame(rows)},
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_addition_asset_category_evidence(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("AT-002d", results[0].rule_id)
        self.assertIn("属性4", results[0].message)

    def test_dt002c_reports_missing_disposal_proceeds_evidence(self):
        rows = [[""] * 21 for _ in range(4)]
        rows[0][18] = "EY Works"
        rows[1][17] = "获得的证据/支持的描述"
        rows[1][18] = "1"
        rows[1][19] = "2"
        rows[1][20] = "3"
        rows[2][5] = "D001"
        rows[2][17] = "报废审批单编号 BF-001"
        rows[2][18] = "Y"
        rows[2][19] = "Y"
        rows[2][20] = "Y"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.02.2 处置测试": pd.DataFrame(rows)},
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_disposal_proceeds_evidence(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DT-002c", results[0].rule_id)

    def test_at003a_reports_cip_transfer_addition(self):
        rows = [[""] * 6 for _ in range(8)]
        rows[0][0] = "原值"
        rows[1][0] = "在建工程转入"
        rows[1][3] = 120
        rows[2][0] = "累计折旧"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.01 Agree SL to GL": pd.DataFrame(rows)},
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_cip_transfer_addition_attributes(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("AT-003a", results[0].rule_id)
        self.assertIn("在建工程转固", results[0].message)

    def test_dp001c_does_not_treat_residual_rate_header_as_residual_review(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.3 折旧政策复核": pd.DataFrame([
                    ["折旧政策", "使用寿命", "残值率"],
                    ["机器设备", "5年", "5%"],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_depreciation_policy_reasonableness_evidence(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-001c", results[0].rule_id)
        self.assertIn("处置残值回顾", results[0].message)

    def test_dp006e_flags_zero_depreciation_exclusion(self):
        tod_rows = [[""] * 18 for _ in range(36)]
        tod_rows[30][1] = "样本数量"
        tod_rows[30][2] = "样本类型"
        tod_rows[30][3] = "固定资产类别"
        tod_rows[30][15] = "差异"
        tod_rows[30][16] = "获得的证据"
        tod_rows[30][17] = "EY Works"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.2 折旧测试TOD": pd.DataFrame(tod_rows),
                "K.03.2 选样输出": pd.DataFrame([
                    ["剔除项", "本期计提折旧为0资产"],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_tod_sampling_exclusions(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-006e", results[0].rule_id)
        self.assertEqual(Severity.HIGH, results[0].severity)

    def test_dp006f_flags_key_item_amount_mismatch(self):
        tod_rows = [[""] * 18 for _ in range(36)]
        tod_rows[0][0] = "关键项金额"
        tod_rows[0][1] = 100
        tod_rows[30][1] = "样本数量"
        tod_rows[30][2] = "样本类型"
        tod_rows[30][3] = "固定资产类别"
        tod_rows[30][15] = "差异"
        tod_rows[30][16] = "获得的证据"
        tod_rows[30][17] = "EY Works"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.2 折旧测试TOD": pd.DataFrame(tod_rows),
                "K.03.2 选样输出": pd.DataFrame([
                    ["关键项金额", 200],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_tod_key_item_amount_match(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-006f", results[0].rule_id)

    def test_dp006g_flags_non_random_sampling_method(self):
        tod_rows = [[""] * 18 for _ in range(36)]
        tod_rows[30][1] = "样本数量"
        tod_rows[30][2] = "样本类型"
        tod_rows[30][3] = "固定资产类别"
        tod_rows[30][15] = "差异"
        tod_rows[30][16] = "获得的证据"
        tod_rows[30][17] = "EY Works"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.2 折旧测试TOD": pd.DataFrame(tod_rows),
                "K.03.2 选样输出": pd.DataFrame([
                    ["双重目的", "否"],
                    ["主要审计问题", "多计"],
                    ["保证水平", "几乎没有"],
                    ["抽样方法", "系统抽样"],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_tod_sampling_method_params(workbook)

        self.assertTrue(any(result.rule_id == "DP-006g" for result in results))
        self.assertTrue(any("抽样方法" in result.message for result in results))

    def test_dp006h_flags_missing_representative_sampling_evaluation(self):
        tod_rows = [[""] * 18 for _ in range(36)]
        tod_rows[30][1] = "样本数量"
        tod_rows[30][2] = "样本类型"
        tod_rows[30][3] = "固定资产类别"
        tod_rows[30][15] = "差异"
        tod_rows[30][16] = "获得的证据"
        tod_rows[30][17] = "EY Works"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.2 折旧测试TOD": pd.DataFrame(tod_rows),
                "K.03.2 选样输出": pd.DataFrame([
                    ["抽样方法", "随机抽样"],
                    ["样本数量", 10],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_tod_representative_sampling_evaluation(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-006h", results[0].rule_id)

    def test_dp007a_reports_missing_policy_evidence(self):
        rows = [[""] * 20 for _ in range(36)]
        rows[30][15] = "差异"
        rows[30][16] = "获得的证据"
        rows[30][17] = "EY Works"
        rows[31][17] = "1"
        rows[31][18] = "2"
        rows[32][16] = "合同发票"
        rows[32][17] = "Y"
        rows[32][18] = "Y"
        workbook = AuditWorkbook(file_path="test.xlsx", sheets={"K.03.2 折旧测试TOD": pd.DataFrame(rows)})
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_tod_policy_evidence(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-007a", results[0].rule_id)

    def test_dp007c_reports_difference_without_follow_up(self):
        rows = [[""] * 20 for _ in range(36)]
        rows[30][15] = "差异"
        rows[30][16] = "获得的证据"
        rows[30][17] = "EY Works"
        rows[31][17] = "1"
        rows[31][18] = "2"
        rows[32][15] = 10
        rows[32][16] = "已重算"
        rows[32][17] = "Y"
        rows[32][18] = "Y"
        workbook = AuditWorkbook(file_path="test.xlsx", sheets={"K.03.2 折旧测试TOD": pd.DataFrame(rows)})
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_tod_difference_follow_up(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-007c", results[0].rule_id)

    def test_dp007d_reports_short_evidence_description(self):
        rows = [[""] * 20 for _ in range(36)]
        rows[30][16] = "获得的证据"
        rows[30][17] = "EY Works"
        rows[31][17] = "1"
        rows[31][18] = "2"
        rows[32][16] = "OK"
        rows[32][17] = "Y"
        rows[32][18] = "Y"
        workbook = AuditWorkbook(file_path="test.xlsx", sheets={"K.03.2 折旧测试TOD": pd.DataFrame(rows)})
        rule = CompletenessRules(enable_ai=False)

        results = rule._check_tod_sample_evidence_completeness(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-007d", results[0].rule_id)

    def test_dp006_matches_tod_samples_by_asset_id_header(self):
        tod_rows = [[""] * 18 for _ in range(46)]
        tod_rows[30][1] = "样本数量"
        tod_rows[30][2] = "样本类型"
        tod_rows[30][3] = "固定资产类别"
        tod_rows[30][15] = "差异"
        tod_rows[30][16] = "获得的证据"
        tod_rows[30][17] = "EY Works"
        tod_rows[42][1] = "固定资产编号"
        tod_rows[42][2] = "样本类型"
        tod_rows[42][3] = "资产名称"
        tod_rows[43][1] = "A001"
        tod_rows[43][2] = "代表性样本"
        tod_rows[43][3] = "设备A"
        tod_rows[44][1] = "A999"
        tod_rows[44][2] = "代表性样本"
        tod_rows[44][3] = "设备B"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.2 折旧测试TOD": pd.DataFrame(tod_rows),
                "K.03.2 选样输出": pd.DataFrame([
                    ["固定资产编号", "资产名称", "折旧额"],
                    ["A001", "设备A", 100],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_tod_sample_process(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-006", results[0].rule_id)
        self.assertIn("A999", results[0].message)

    def test_dp006_does_not_treat_non_id_columns_as_sample_ids(self):
        tod_rows = [[""] * 18 for _ in range(45)]
        tod_rows[30][1] = "样本数量"
        tod_rows[30][2] = "样本类型"
        tod_rows[30][3] = "固定资产类别"
        tod_rows[30][15] = "差异"
        tod_rows[30][16] = "获得的证据"
        tod_rows[30][17] = "EY Works"
        tod_rows[42][1] = "固定资产编号"
        tod_rows[42][2] = "样本类型"
        tod_rows[42][3] = "固定资产类别"
        tod_rows[42][15] = "差异"
        tod_rows[42][16] = "获得的证据"
        tod_rows[42][17] = "EY Works"
        tod_rows[43][1] = "A001"
        tod_rows[43][2] = "代表性样本"
        tod_rows[43][3] = "机器设备"
        tod_rows[43][15] = 0
        tod_rows[43][16] = "2026-03-31"
        tod_rows[43][17] = "Y"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.2 折旧测试TOD": pd.DataFrame(tod_rows),
                "K.03.2 选样输出": pd.DataFrame([
                    ["固定资产编号", "资产名称", "折旧额"],
                    ["A001", "设备A", 100],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)

        results = rule._check_tod_sample_process(workbook)

        self.assertEqual([], results)

    def test_dp007_without_tod_records_business_skip_reason(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.2 By Item": pd.DataFrame([["固定资产编号", "原值"]]),
            },
        )
        rule = CompletenessRules(enable_ai=False)
        rule.reset_runtime_state()

        results = rule._check_dp007(workbook)

        self.assertEqual([], results)
        self.assertIn("DP-007", rule.runtime_skips)
        self.assertIn("不存在TOD折旧测试", rule.runtime_skips["DP-007"])

    def test_dp007_ai_error_is_marked_skipped_not_business_issue(self):
        rows = [[""] * 20 for _ in range(36)]
        rows[30][1] = "样本数量"
        rows[30][2] = "样本类型"
        rows[30][3] = "固定资产类别"
        rows[30][15] = "差异"
        rows[30][16] = "获得的证据"
        rows[30][17] = "EY Works"
        rows[31][17] = "1"
        rows[32][17] = "N"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.03.2 折旧测试TOD": pd.DataFrame(rows)},
        )
        rule = CompletenessRules(enable_ai=True)
        rule.ai_client = ErrorAIClient()
        rule.reset_runtime_state()

        results = rule._check_dp007(workbook)

        self.assertEqual([], results)
        self.assertIn("DP-007", rule.runtime_skips)
        self.assertIn("AI调用失败，未完成AI判断", rule.runtime_skips["DP-007"])

    def test_gl004_missing_table4_is_marked_skipped(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([["名义金额", 100]]),
                "K.01 Agree SL to GL": pd.DataFrame([["折旧费用核对", 0]]),
            },
        )
        rule = ConsistencyRules()
        rule.reset_runtime_state()

        results = rule._check_depreciation_vs_pnl(workbook)

        self.assertEqual([], results)
        self.assertIn("GL-004", rule.runtime_skips)
        self.assertIn("未识别到表4", rule.runtime_skips["GL-004"])

    def test_dp001_unrecognized_policy_header_is_marked_skipped(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.3 折旧政策复核": pd.DataFrame([
                    ["资产类别", "使用寿命", "残值率"],
                    ["机器设备", "5年", "5%"],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)
        rule.reset_runtime_state()

        results = rule._check_depreciation_policy_change(workbook)

        self.assertEqual([], results)
        self.assertIn("DP-001", rule.runtime_skips)
        self.assertIn("未识别到可核对的折旧政策差异表头", rule.runtime_skips["DP-001"])

    def test_sp003e_reports_missing_representative_sampling_evaluation(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.02.1 新增测试": pd.DataFrame([
                    ["选择关键项目的理由", "按金额超过TT选取关键项"],
                    ["样本类型", "资产名称"],
                    ["关键项", "设备A"],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)
        rule.reset_runtime_state()

        results = rule._check_addition_representative_sampling_evaluation(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-003e", results[0].rule_id)
        self.assertIn("代表性样本", results[0].message)

    def test_sp004e_reports_missing_disposal_completeness_categories(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.00 Lead Sheet": pd.DataFrame([
                    ["处置完整性", "已询问管理层"],
                ]),
                "K.02.2 处置测试": pd.DataFrame([
                    ["处置测试", "已询问管理层"],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)
        rule.reset_runtime_state()

        results = rule._check_disposal_completeness_evaluation(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("SP-004e", results[0].rule_id)
        self.assertIn("收入分析", results[0].message)
        self.assertIn("预期核对", results[0].message)

    def test_dt003_reports_other_nature_reduction(self):
        rows = [[""] * 6 for _ in range(8)]
        rows[0][0] = "原值"
        rows[1][0] = "债务重组减少"
        rows[1][3] = -120
        rows[2][0] = "累计折旧"
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={"K.01 Agree SL to GL": pd.DataFrame(rows)},
        )
        rule = CompletenessRules(enable_ai=False)
        rule.reset_runtime_state()

        findings = rule._find_special_nature_changes(workbook)
        results = rule._check_disposal_special_nature(findings)

        self.assertEqual(1, len(results))
        self.assertEqual("DT-003", results[0].rule_id)
        self.assertIn("债务重组减少", results[0].message)

    def test_dp001b_reports_policy_change_without_approval_evidence(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.3 折旧政策复核": pd.DataFrame([
                    ["折旧政策", "本期使用年限", "本期残值率", "本期年折旧率", "上期使用年限", "上期残值率", "上期年折旧率", "使用年限差异", "残值率差异", "差异说明"],
                    ["机器设备", "5年", "5%", "", "10年", "5%", "", False, True, "因政策变更未来适用"],
                ]),
            },
        )
        rule = CompletenessRules(enable_ai=False)
        rule.reset_runtime_state()

        results = rule._check_depreciation_policy_change_supporting_files(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-001b", results[0].rule_id)
        self.assertIn("审批决议", results[0].message)

    def test_dp002d_reports_policy_mismatch_without_resolution_record(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.3 折旧政策复核": pd.DataFrame([
                    ["折旧政策", "使用寿命"],
                    ["机器设备", "5年"],
                ]),
                "FA list": pd.DataFrame([
                    ["资产类别", "使用年限"],
                    ["机器设备", 10],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)
        rule.reset_runtime_state()

        results = rule._check_policy_inconsistency_resolution(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-002d", results[0].rule_id)
        self.assertIn("未识别到原因调查", results[0].message)

    def test_dp002e_reports_external_disclosure_dependency(self):
        workbook = AuditWorkbook(
            file_path="test.xlsx",
            sheets={
                "K.03.3 折旧政策复核": pd.DataFrame([
                    ["折旧政策", "使用寿命"],
                    ["机器设备", "5年"],
                ]),
            },
        )
        rule = LogicRules(enable_ai=False)
        rule.reset_runtime_state()

        results = rule._check_policy_report_disclosure_consistency(workbook)

        self.assertEqual(1, len(results))
        self.assertEqual("DP-002e", results[0].rule_id)
        self.assertIn("财报/披露底稿", results[0].message)


if __name__ == "__main__":
    unittest.main()
