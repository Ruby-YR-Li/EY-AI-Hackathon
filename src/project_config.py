"""项目基础信息、重要性水平和科目认定配置的 Excel 模板与校验。"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName


SUBJECTS = [
    ("fixed_assets", "固定资产"),
    ("intangible_assets", "无形资产"),
    ("construction_in_progress", "在建工程"),
    ("long_term_deferred_expenses", "长期待摊费用"),
    ("cash", "货币资金"),
    ("management_sales_expenses", "管理费用&销售费用"),
    ("finance_expenses", "财务费用"),
]
SUBJECT_NAME_TO_CODE = {name: code for code, name in SUBJECTS}
SUBJECT_CODE_TO_NAME = dict(SUBJECTS)

ASSERTIONS = [
    ("C", "完整性（C）"),
    ("EO", "存在性/发生（E/O）"),
    ("VM", "计价/计量（V/M）"),
    ("RO", "权利和义务（R&O）"),
    ("PD", "列报和披露（P&D）"),
]
CRA_OPTIONS = ["minimal", "low", "moderate", "high", "N/A"]
CRA_ALIASES = {
    "minimal": "minimal",
    "最低": "minimal",
    "low": "low",
    "低": "low",
    "moderate": "moderate",
    "medium": "moderate",
    "中": "moderate",
    "high": "high",
    "高": "high",
    "n/a": "N/A",
    "na": "N/A",
    "不适用": "N/A",
}

TEMPLATE_VERSION = "1.0"
BASIC_INFO_SHEET = "项目基础信息"
ASSERTION_SHEET = "科目认定配置"
DICTIONARY_SHEET = "数据字典"


class ProjectConfigValidationError(ValueError):
    """项目配置校验错误。"""


def normalize_cra(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    normalized = CRA_ALIASES.get(text.lower())
    if normalized is None:
        raise ProjectConfigValidationError(
            f'CRA值"{text}"无效，应为：{", ".join(CRA_OPTIONS)}'
        )
    return normalized


def parse_threshold_pct(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        if text.endswith("%"):
            try:
                number = float(text[:-1].strip()) / 100
            except ValueError as exc:
                raise ProjectConfigValidationError(f'Threshold %"{value}"不是有效百分比') from exc
        else:
            try:
                number = float(text)
            except ValueError as exc:
                raise ProjectConfigValidationError(f'Threshold %"{value}"不是有效数值') from exc
    elif isinstance(value, (int, float)):
        number = float(value)
    else:
        raise ProjectConfigValidationError(f'Threshold %"{value}"不是有效数值')
    if number < 0 or number > 1:
        raise ProjectConfigValidationError("Threshold %必须在0%至100%之间")
    return number


def _parse_number(value: Any, label: str) -> float:
    if value is None or str(value).strip() == "":
        raise ProjectConfigValidationError(f"{label}不能为空")
    try:
        return float(str(value).replace(",", "").replace("，", "").strip())
    except ValueError as exc:
        raise ProjectConfigValidationError(f"{label}必须为数字") from exc


def _parse_date(value: Any) -> str:
    if value is None or str(value).strip() == "":
        raise ProjectConfigValidationError("资产负债表日不能为空")
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise ProjectConfigValidationError("资产负债表日格式无效，应为YYYY-MM-DD")


def validate_project_config(
    project_code: str,
    basic_info: Optional[dict],
    materiality: Optional[dict],
    subject_assertions: Optional[dict],
    require_assertions: bool = False,
) -> dict:
    code = str(project_code or "").strip()
    info = dict(basic_info or {})
    mat = dict(materiality or {})
    assertions = dict(subject_assertions or {})

    if not code:
        raise ProjectConfigValidationError("Engagement Code不能为空")

    balance_sheet_date = _parse_date(info.get("balance_sheet_date"))
    pm = _parse_number(mat.get("pm"), "PM")
    te = _parse_number(mat.get("te"), "TE")
    sad = _parse_number(mat.get("sad"), "SAD")
    if pm <= 0:
        raise ProjectConfigValidationError("PM必须大于0")
    if te <= 0 or te > pm:
        raise ProjectConfigValidationError("TE必须大于0且小于或等于PM")
    if sad < 0 or sad > te:
        raise ProjectConfigValidationError("SAD必须大于或等于0且小于或等于TE")

    normalized_assertions: Dict[str, dict] = {}
    errors = []
    for subject_code, assertion_values in assertions.items():
        if subject_code not in SUBJECT_CODE_TO_NAME:
            errors.append(f"不支持的科目代码：{subject_code}")
            continue
        values = assertion_values or {}
        normalized_values = {}
        configured_count = 0
        for assertion_code, assertion_name in ASSERTIONS:
            item = values.get(assertion_code) or {}
            raw_cra = item.get("cra")
            raw_threshold = item.get("threshold_pct")
            if raw_cra in (None, "") and raw_threshold in (None, ""):
                continue
            configured_count += 1
            try:
                cra = normalize_cra(raw_cra)
                if cra is None:
                    raise ProjectConfigValidationError("CRA不能为空")
                threshold = parse_threshold_pct(raw_threshold)
                if cra == "N/A" and threshold is not None:
                    raise ProjectConfigValidationError("CRA为N/A时Threshold %必须为空")
                if cra != "N/A" and threshold is None:
                    raise ProjectConfigValidationError("CRA非N/A时Threshold %不能为空")
                normalized_values[assertion_code] = {
                    "cra": cra,
                    "threshold_pct": threshold,
                }
            except ProjectConfigValidationError as exc:
                errors.append(
                    f"{SUBJECT_CODE_TO_NAME[subject_code]}-{assertion_name}：{exc}"
                )
        if configured_count:
            missing = [
                name for code_key, name in ASSERTIONS if code_key not in normalized_values
            ]
            if missing:
                errors.append(
                    f"{SUBJECT_CODE_TO_NAME[subject_code]}：五项认定必须填写完整，缺少{', '.join(missing)}"
                )
            else:
                normalized_assertions[subject_code] = normalized_values

    if require_assertions and not normalized_assertions:
        errors.append("至少需要配置一个科目的五项认定")
    if errors:
        raise ProjectConfigValidationError("；".join(errors))

    info.pop("engagement_name", None)
    info.pop("project_name", None)
    info["balance_sheet_date"] = balance_sheet_date
    return {
        "project_code": code,
        "basic_info": info,
        "materiality": {"pm": pm, "te": te, "sad": sad},
        "subject_assertions": normalized_assertions,
    }


def build_project_config_template() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = BASIC_INFO_SHEET

    title_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    title_font = Font(color="FFFFFF", bold=True)

    ws["A1"] = "项目基础信息配置表"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws.merge_cells("A1:C1")
    ws["A2"] = "请填写B列；金额必须为数字，日期格式为YYYY-MM-DD。"
    ws.merge_cells("A2:C2")
    ws["A2"].alignment = Alignment(wrap_text=True)
    fields = [
        ("Engagement Code", "", "必填，项目唯一编号"),
        ("资产负债表日", "", "YYYY-MM-DD"),
        ("PM", "", "整体重要性水平，PM > 0"),
        ("TE", "", "0 < TE <= PM"),
        ("SAD", "", "0 <= SAD <= TE"),
    ]
    for row_idx, (label, value, note) in enumerate(fields, start=4):
        ws.cell(row_idx, 1, label)
        ws.cell(row_idx, 2, value)
        ws.cell(row_idx, 3, note)
        ws.cell(row_idx, 1).fill = section_fill
        ws.cell(row_idx, 1).font = Font(bold=True)
    ws["B5"].number_format = "yyyy-mm-dd"
    for cell in ("B6", "B7", "B8"):
        ws[cell].number_format = '#,##0.00'
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 42
    ws.freeze_panes = "A4"

    assertion_ws = wb.create_sheet(ASSERTION_SHEET)
    headers = ["科目"]
    for _, assertion_name in ASSERTIONS:
        headers.extend([f"{assertion_name} CRA", f"{assertion_name} Threshold %"])
    for col_idx, header in enumerate(headers, start=1):
        cell = assertion_ws.cell(1, col_idx, header)
        cell.fill = title_fill
        cell.font = title_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    assertion_ws.row_dimensions[1].height = 42
    assertion_ws.freeze_panes = "A2"
    assertion_ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}31"

    for row_idx, (_, subject_name) in enumerate(SUBJECTS, start=2):
        assertion_ws.cell(row_idx, 1, subject_name)
        for col_idx in range(3, len(headers) + 1, 2):
            assertion_ws.cell(row_idx, col_idx).number_format = "0.00%"
    for row_idx in range(2 + len(SUBJECTS), 32):
        for col_idx in range(3, len(headers) + 1, 2):
            assertion_ws.cell(row_idx, col_idx).number_format = "0.00%"

    assertion_ws.column_dimensions["A"].width = 24
    for col_idx in range(2, len(headers) + 1):
        assertion_ws.column_dimensions[get_column_letter(col_idx)].width = 19

    dictionary_ws = wb.create_sheet(DICTIONARY_SHEET)
    dictionary_ws.append(["科目", "CRA", "模板版本"])
    max_rows = max(len(SUBJECTS), len(CRA_OPTIONS))
    for idx in range(max_rows):
        dictionary_ws.append([
            SUBJECTS[idx][1] if idx < len(SUBJECTS) else None,
            CRA_OPTIONS[idx] if idx < len(CRA_OPTIONS) else None,
            TEMPLATE_VERSION if idx == 0 else None,
        ])
    wb.defined_names.add(DefinedName(
        "SubjectOptions",
        attr_text=f"'{DICTIONARY_SHEET}'!$A$2:$A${len(SUBJECTS) + 1}",
    ))
    wb.defined_names.add(DefinedName(
        "CRAOptions",
        attr_text=f"'{DICTIONARY_SHEET}'!$B$2:$B${len(CRA_OPTIONS) + 1}",
    ))
    subject_validation = DataValidation(type="list", formula1="=SubjectOptions")
    cra_validation = DataValidation(type="list", formula1="=CRAOptions")
    assertion_ws.add_data_validation(subject_validation)
    assertion_ws.add_data_validation(cra_validation)
    subject_validation.add("A2:A31")
    for col_idx in range(2, len(headers) + 1, 2):
        cra_validation.add(
            f"{get_column_letter(col_idx)}2:{get_column_letter(col_idx)}31"
        )
    dictionary_ws.sheet_state = "hidden"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def parse_project_config_workbook(file_obj) -> dict:
    try:
        wb = load_workbook(file_obj, data_only=True)
    except Exception as exc:
        raise ProjectConfigValidationError(f"无法读取Excel配置表：{exc}") from exc
    try:
        if BASIC_INFO_SHEET not in wb.sheetnames or ASSERTION_SHEET not in wb.sheetnames:
            raise ProjectConfigValidationError(
                f'配置表必须包含"{BASIC_INFO_SHEET}"和"{ASSERTION_SHEET}"Sheet'
            )

        basic_ws = wb[BASIC_INFO_SHEET]
        field_values = {}
        for row_idx in range(1, basic_ws.max_row + 1):
            label = basic_ws.cell(row_idx, 1).value
            if label is not None:
                field_values[str(label).strip()] = basic_ws.cell(row_idx, 2).value
        basic_info = {
            "balance_sheet_date": field_values.get("资产负债表日"),
        }
        materiality = {
            "pm": field_values.get("PM"),
            "te": field_values.get("TE"),
            "sad": field_values.get("SAD"),
        }

        assertion_ws = wb[ASSERTION_SHEET]
        subject_assertions = {}
        duplicate_subjects = []
        for row_idx in range(2, assertion_ws.max_row + 1):
            subject_name = assertion_ws.cell(row_idx, 1).value
            if subject_name is None or not str(subject_name).strip():
                continue
            subject_name = str(subject_name).strip()
            subject_code = SUBJECT_NAME_TO_CODE.get(subject_name)
            if subject_code is None:
                raise ProjectConfigValidationError(f'第{row_idx}行科目"{subject_name}"不在下拉清单中')
            if subject_code in subject_assertions:
                duplicate_subjects.append(subject_name)
                continue
            assertion_values = {}
            for assertion_idx, (assertion_code, _) in enumerate(ASSERTIONS):
                cra_col = 2 + assertion_idx * 2
                threshold_col = cra_col + 1
                assertion_values[assertion_code] = {
                    "cra": assertion_ws.cell(row_idx, cra_col).value,
                    "threshold_pct": assertion_ws.cell(row_idx, threshold_col).value,
                }
            subject_assertions[subject_code] = assertion_values
        if duplicate_subjects:
            raise ProjectConfigValidationError(
                f"科目不能重复：{', '.join(sorted(set(duplicate_subjects)))}"
            )

        return validate_project_config(
            field_values.get("Engagement Code"),
            basic_info,
            materiality,
            subject_assertions,
            require_assertions=True,
        )
    finally:
        wb.close()
