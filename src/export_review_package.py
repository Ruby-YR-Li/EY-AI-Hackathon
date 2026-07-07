"""Export a traceable review package for expense workpaper findings."""

from __future__ import annotations

import html
import re
import shutil
from dataclasses import asdict
from pathlib import Path
from typing import Iterable

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from expense_review_engine import DEFAULT_FILES, ReviewNote, ReviewSummary, notes_to_csv, select_workpapers


OUTPUT_DIR = Path(__file__).resolve().parents[1] / "outputs"
NOTES_SHEET = "AI Review Notes"

HIGH_FILL = PatternFill("solid", fgColor="FFF2CC")
MEDIUM_FILL = PatternFill("solid", fgColor="FFFCE5")
HIGH_SIDE = Side(style="medium", color="C00000")
MEDIUM_SIDE = Side(style="thin", color="BF9000")
HEADER_FILL = PatternFill("solid", fgColor="111111")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _esc(value: object) -> str:
    return html.escape("" if value is None else str(value), quote=True)


def _row(note: ReviewNote) -> dict[str, str]:
    return asdict(note)


def _rows(notes: Iterable[ReviewNote]) -> list[dict[str, str]]:
    return [_row(note) for note in notes]


def _safe_stem(path: Path) -> str:
    stem = re.sub(r"[^\w\u4e00-\u9fff]+", "_", path.stem, flags=re.UNICODE)
    stem = re.sub(r"_+", "_", stem).strip("_")
    return stem or "workpaper"


def _target_cell(location: str) -> str | None:
    """Return the first A1 coordinate from a location string."""
    if not location:
        return None
    match = re.search(r"\$?([A-Z]{1,3})\$?([0-9]{1,7})", location)
    if not match:
        return None
    return f"{match.group(1)}{match.group(2)}"


def export_review_notes_xlsx(
    notes: list[ReviewNote],
    summary: ReviewSummary,
    output_path: Path,
) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Review Notes"

    headers = [
        "risk_level",
        "file",
        "sheet",
        "location",
        "issue_type",
        "review_note",
        "suggested_action",
        "source",
        "judge_method",
        "confidence",
        "evidence_summary",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    order = {"High": 0, "Medium": 1, "Low": 2}
    for note in sorted(notes, key=lambda n: (order.get(n.risk_level, 9), n.file, n.sheet, n.location)):
        data = _row(note)
        ws.append([data.get(h, "") for h in headers])

    widths = {
        "A": 12,
        "B": 34,
        "C": 28,
        "D": 14,
        "E": 24,
        "F": 60,
        "G": 60,
        "H": 32,
        "I": 16,
        "J": 12,
        "K": 60,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")

    info = wb.create_sheet("Summary")
    info.append(["Metric", "Value"])
    for cell in info[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for key, value in asdict(summary).items():
        info.append([key, value])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 18

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def export_annotated_workbook(
    notes: list[ReviewNote],
    source_path: Path,
    output_path: Path,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, output_path)

    wb = openpyxl.load_workbook(output_path)
    if NOTES_SHEET in wb.sheetnames:
        del wb[NOTES_SHEET]
    ws_notes = wb.create_sheet(NOTES_SHEET, 0)
    headers = [
        "risk_level",
        "file",
        "sheet",
        "location",
        "issue_type",
        "review_note",
        "suggested_action",
        "source",
        "judge_method",
        "confidence",
        "evidence_summary",
    ]
    ws_notes.append(headers)
    for cell in ws_notes[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for idx, note in enumerate(notes, start=2):
        data = _row(note)
        ws_notes.append([data.get(h, "") for h in headers])

        if note.file != source_path.name or note.sheet not in wb.sheetnames:
            continue
        coord = _target_cell(note.location)
        if not coord:
            continue
        ws = wb[note.sheet]
        cell = ws[coord]
        comment_body = (
            f"[{note.risk_level}] {note.issue_type}\n"
            f"{note.review_note}\n\n"
            f"建议：{note.suggested_action}\n"
            f"来源：{note.source}\n"
            f"判断方式：{note.judge_method} {note.confidence}\n"
            f"详见 {NOTES_SHEET}!A{idx}:K{idx}"
        )
        if cell.comment:
            comment_body = f"{cell.comment.text}\n\n--- AI Review ---\n{comment_body}"
        cell.comment = Comment(comment_body, "AI Review Assistant")
        cell.fill = HIGH_FILL if note.risk_level == "High" else MEDIUM_FILL
        side = HIGH_SIDE if note.risk_level == "High" else MEDIUM_SIDE
        cell.border = Border(left=side, right=side, top=side, bottom=side)

    for col_idx, width in enumerate([12, 34, 28, 14, 24, 60, 60, 32, 16, 12, 60], start=1):
        ws_notes.column_dimensions[get_column_letter(col_idx)].width = width
    ws_notes.freeze_panes = "A2"
    ws_notes.auto_filter.ref = ws_notes.dimensions
    for row in ws_notes.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")

    # ── AI Review 执行日志 sheet ──
    _add_execution_log_sheet(wb, notes)

    wb.save(output_path)
    return output_path


EXEC_LOG_SHEET = "AI Review 执行日志"


def _add_execution_log_sheet(wb, notes: list[ReviewNote]) -> None:
    """在标注底稿中添加质检点执行日志 sheet。"""
    if EXEC_LOG_SHEET in wb.sheetnames:
        del wb[EXEC_LOG_SHEET]
    ws = wb.create_sheet(EXEC_LOG_SHEET)

    # Aggregate notes by issue_type
    from collections import Counter
    type_counts: dict[str, Counter] = {}
    type_examples: dict[str, str] = {}
    for n in notes:
        if n.issue_type not in type_counts:
            type_counts[n.issue_type] = Counter()
            type_examples[n.issue_type] = n.review_note[:80]
        type_counts[n.issue_type][n.risk_level] += 1

    headers = ["规则名称", "High", "Medium", "Low", "总计", "示例"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    order_rank = {"High": 1, "Medium": 2, "Low": 3}
    sorted_types = sorted(
        type_counts.items(),
        key=lambda item: (
            -item[1].get("High", 0),
            -item[1].get("Medium", 0),
            item[0],
        )
    )
    for issue_type, counts in sorted_types:
        total = sum(counts.values())
        ws.append([
            issue_type,
            counts.get("High", 0),
            counts.get("Medium", 0),
            counts.get("Low", 0),
            total,
            type_examples.get(issue_type, ""),
        ])
        # Color High count red
        row_idx = ws.max_row
        high_cell = ws.cell(row_idx, 2)
        if high_cell.value and high_cell.value > 0:
            high_cell.font = Font(color="B42318", bold=True)

    widths = {"A": 32, "B": 10, "C": 10, "D": 10, "E": 10, "F": 60}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")


def export_html_report(
    notes: list[ReviewNote],
    summary: ReviewSummary,
    output_path: Path,
    *,
    review_notes_xlsx: Path,
    annotated_workbook: Path,
) -> Path:
    rows = ""
    order = {"High": 0, "Medium": 1, "Low": 2}
    for idx, note in enumerate(sorted(notes, key=lambda n: (order.get(n.risk_level, 9), n.file, n.sheet, n.location)), start=1):
        rows += f"""
        <tr class="risk-{_esc(note.risk_level).lower()}">
          <td>{idx}</td>
          <td><span class="badge {_esc(note.risk_level).lower()}">{_esc(note.risk_level)}</span></td>
          <td>{_esc(note.file)}</td>
          <td>{_esc(note.sheet)}</td>
          <td><code>{_esc(note.location)}</code></td>
          <td>{_esc(note.issue_type)}</td>
          <td>{_esc(note.review_note)}</td>
          <td>{_esc(note.suggested_action)}</td>
          <td>{_esc(note.source)}</td>
          <td>{_esc(note.judge_method)}</td>
          <td>{_esc(note.confidence)}</td>
          <td>{_esc(note.evidence_summary)}</td>
        </tr>
        """

    html_doc = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <title>费用底稿 AI Review 报告</title>
  <style>
    body {{ margin: 0; background: #f5f5f5; color: #111; font-family: "Segoe UI", "Microsoft YaHei", Arial, sans-serif; }}
    .topbar {{ background: #111; color: #fff; border-left: 10px solid #ffe600; padding: 20px 28px; }}
    .topbar h1 {{ margin: 0; font-size: 26px; }}
    .topbar p {{ margin: 8px 0 0; color: #d0d0d0; }}
    main {{ padding: 22px 28px; }}
    .cards {{ display: grid; grid-template-columns: repeat(6, minmax(120px, 1fr)); gap: 12px; margin-bottom: 18px; }}
    .card {{ background: #fff; border: 1px solid #ddd; border-left: 5px solid #ffe600; border-radius: 6px; padding: 14px; }}
    .label {{ color: #667085; font-size: 13px; }}
    .value {{ font-size: 28px; font-weight: 750; margin-top: 4px; }}
    .links a {{ display: inline-block; margin: 0 10px 12px 0; padding: 9px 12px; border-radius: 4px; background: #111; color: #fff; text-decoration: none; }}
    table {{ width: 100%; border-collapse: collapse; background: #fff; font-size: 13px; }}
    th, td {{ border: 1px solid #ddd; padding: 8px 10px; vertical-align: top; }}
    th {{ background: #222; color: #fff; position: sticky; top: 0; }}
    .badge {{ padding: 3px 8px; border-radius: 999px; color: #fff; font-weight: 700; font-size: 12px; }}
    .badge.high {{ background: #b42318; }}
    .badge.medium {{ background: #b54708; }}
    .badge.low {{ background: #175cd3; }}
    tr.risk-high td {{ background: #fff7f5; }}
    tr.risk-medium td {{ background: #fffaf0; }}
    code {{ font-family: Consolas, monospace; }}
  </style>
</head>
<body>
  <div class="topbar">
    <h1>费用底稿 AI Review 报告</h1>
    <p>基于 SOP、审计程序要求和 VC&VD 底稿自动生成 Review Notes，并输出可追溯附件。</p>
  </div>
  <main>
    <section class="cards">
      <div class="card"><div class="label">Review Notes</div><div class="value">{summary.total_notes}</div></div>
      <div class="card"><div class="label">High</div><div class="value">{summary.high_count}</div></div>
      <div class="card"><div class="label">Medium</div><div class="value">{summary.medium_count}</div></div>
      <div class="card"><div class="label">Low</div><div class="value">{summary.low_count}</div></div>
      <div class="card"><div class="label">涉及 Sheet</div><div class="value">{summary.sheets_impacted}</div></div>
      <div class="card"><div class="label">规则数</div><div class="value">{summary.rules_run}</div></div>
    </section>
    <section class="links">
      <a href="{_esc(review_notes_xlsx.name)}">下载 Review Notes Excel</a>
      <a href="{_esc(annotated_workbook.name)}">下载标注底稿副本</a>
      <a href="review_notes.csv">下载 CSV 备用</a>
    </section>
    <table>
      <thead>
        <tr>
          <th>#</th><th>风险</th><th>文件</th><th>Sheet</th><th>位置</th>
          <th>问题类别</th><th>Review Note</th><th>建议动作</th><th>来源</th>
          <th>判断方式</th><th>置信度</th><th>证据摘要</th>
        </tr>
      </thead>
      <tbody>{rows or '<tr><td colspan="12">未发现 Review Notes</td></tr>'}</tbody>
    </table>
  </main>
</body>
</html>
"""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(html_doc, encoding="utf-8")
    return output_path


def export_review_package(
    notes: list[ReviewNote],
    summary: ReviewSummary,
    *,
    files: dict[str, Path] | None = None,
    output_dir: Path | None = None,
) -> dict[str, Path]:
    paths = dict(DEFAULT_FILES)
    if files:
        for key, value in files.items():
            if value is None and key in {"trial_balance", "general_ledger"}:
                paths[key] = None
                continue
            if not value:
                continue
            if key == "workpapers":
                paths[key] = [Path(p) for p in value]  # type: ignore[arg-type]
            else:
                paths[key] = Path(value)  # type: ignore[arg-type]
    out = output_dir or OUTPUT_DIR
    out.mkdir(parents=True, exist_ok=True)

    main_workpaper = select_workpapers(paths).main_workpaper
    notes_xlsx = out / "review_notes.xlsx"
    annotated = out / f"{_safe_stem(main_workpaper)}_AI_review_annotated.xlsx"
    html_report = out / "review_report.html"
    csv_path = out / "review_notes.csv"

    export_review_notes_xlsx(notes, summary, notes_xlsx)
    export_annotated_workbook(notes, main_workpaper, annotated)
    csv_path.write_bytes(notes_to_csv(notes))
    export_html_report(
        notes,
        summary,
        html_report,
        review_notes_xlsx=notes_xlsx,
        annotated_workbook=annotated,
    )
    return {
        "html_report": html_report,
        "review_notes_xlsx": notes_xlsx,
        "annotated_workbook": annotated,
        "review_notes_csv": csv_path,
    }
