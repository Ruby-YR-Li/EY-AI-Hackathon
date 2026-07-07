"""
审计底稿质检Agent - 底稿解析模块
负责读取和解析Excel格式的审计底稿，支持动态定位单元格
"""

import re
import zipfile
import xml.etree.ElementTree as ET
import pandas as pd
from typing import Dict, List, Optional, Tuple

from .models import AuditWorkbook


class WorkbookParser:
    """底稿解析器"""

    # 费用底稿（销售费用&管理费用）的标准Sheet名称
    STANDARD_SHEETS = [
        "汇总",
        "Uexp.00 Lead",
        "VC.00 销售费用BKD",
        "VD.00 管理费用BKD",
        "VC&VD.01.4 截止性测试",
        "VC&VD.01.2 TOD 预审+剩余期间",
        "SkywindSettingSheet",
    ]

    def __init__(self):
        self.evidence_recorder = None

    def _record_evidence(self, sheet: str, row: int = None, col: int = None, text: str = ""):
        if not self.evidence_recorder:
            return
        try:
            self.evidence_recorder(sheet, row, col, text)
        except Exception:
            pass

    # ── 合并单元格辅助 ──────────────────────────────────────

    @staticmethod
    def _col_letter_to_index(col_str: str) -> int:
        """将Excel列字母转为0-based列索引，如 'A'->0, 'D'->3, 'AA'->26"""
        idx = 0
        for ch in col_str:
            idx = idx * 26 + (ord(ch) - ord('A') + 1)
        return idx - 1

    @staticmethod
    def _parse_range_ref(range_ref: str) -> Tuple[int, int, int, int]:
        """解析 'D4:G4' 为 (row_start, row_end, col_start, col_end)，0-based"""
        match = re.match(r'([A-Z]+)(\d+):([A-Z]+)(\d+)', range_ref.upper())
        if not match:
            return 0, 0, 0, 0
        c1, r1, c2, r2 = match.groups()
        col_start = WorkbookParser._col_letter_to_index(c1)
        col_end = WorkbookParser._col_letter_to_index(c2)
        row_start = int(r1) - 1
        row_end = int(r2) - 1
        return row_start, row_end, col_start, col_end

    @staticmethod
    def _parse_drawing_display_name(display_name: str) -> Optional[str]:
        """Attempt to parse a display name from a sheet relationship to find the real file name."""
        pass  # placeholder

    @staticmethod
    def _get_merged_cells_from_xlsx(file_path: str) -> Dict[str, Dict[Tuple[int, int], Tuple[int, int]]]:
        """
        直接从XLSX内部XML读取合并单元格信息（不依赖openpyxl完整加载）。
        返回 { sheet_name: {(row, col): (top_left_row, top_left_col)} }
        """
        NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        NS_REL = {'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
        merged_map = {}

        try:
            with zipfile.ZipFile(file_path, 'r') as z:
                # 1. 读取workbook.xml -> sheet名 + rId映射
                wb_tree = ET.parse(z.open('xl/workbook.xml'))
                wb_root = wb_tree.getroot()

                # sheet名 by rId
                sheet_name_by_rId = {}
                # 收集隐藏sheet
                hidden_sheets_from_xml = set()

                for sheet_elem in wb_root.findall('.//s:sheet', NS):
                    name = sheet_elem.get('name')
                    rId = sheet_elem.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                    state = sheet_elem.get('state', 'visible')
                    if name and rId:
                        sheet_name_by_rId[rId] = name
                        if state != 'visible':
                            hidden_sheets_from_xml.add(name)

                # 2. 读取workbook.xml.rels -> rId映射到worksheet路径
                rels_tree = ET.parse(z.open('xl/_rels/workbook.xml.rels'))
                rels_root = rels_tree.getroot()
                # 注意rels有自己的命名空间
                REL_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'
                rel_path_by_rId = {}
                for rel_elem in rels_root:
                    rId = rel_elem.get('Id')
                    target = rel_elem.get('Target')
                    if rId and target:
                        rel_path_by_rId[rId] = target

                # 3. 遍历每个sheet，读取其XML中的mergeCells
                for rId, sheet_name in sheet_name_by_rId.items():
                    if sheet_name in hidden_sheets_from_xml:
                        continue

                    rel_path = rel_path_by_rId.get(rId)
                    if not rel_path:
                        continue

                    # 路径可能是相对路径，如 'worksheets/sheet1.xml'
                    ws_path = f'xl/{rel_path}' if not rel_path.startswith('xl/') else rel_path

                    try:
                        ws_tree = ET.parse(z.open(ws_path))
                        ws_root = ws_tree.getroot()
                    except (KeyError, ET.ParseError):
                        continue

                    sheet_merged = {}
                    merge_cells_elem = ws_root.find('s:mergeCells', NS)
                    if merge_cells_elem is not None:
                        for mc in merge_cells_elem.findall('s:mergeCell', NS):
                            ref = mc.get('ref')
                            if not ref:
                                continue
                            r0, r1, c0, c1 = WorkbookParser._parse_range_ref(ref)
                            # 非左上角的单元格映射到左上角
                            for r in range(r0, r1 + 1):
                                for c in range(c0, c1 + 1):
                                    if (r, c) != (r0, c0):
                                        sheet_merged[(r, c)] = (r0, c0)

                    if sheet_merged:
                        merged_map[sheet_name] = sheet_merged

            return merged_map

        except Exception as e:
            print(f"读取合并单元格信息失败（将跳过合并单元格处理）: {e}")
            return {}

    def parse(self, file_path: str) -> AuditWorkbook:
        """解析Excel底稿文件，跳过隐藏的Sheet，处理合并单元格"""
        workbook = AuditWorkbook(file_path=file_path)

        # 使用openpyxl读取Sheet可见性信息（read_only模式，安全快速）
        hidden_sheets = set()
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True, data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                if sheet.sheet_state != 'visible':
                    hidden_sheets.add(sheet_name)
                    print(f"跳过隐藏Sheet: {sheet_name}")
            wb.close()
        except Exception as e:
            print(f"读取Sheet可见性信息失败（将检查所有Sheet）: {e}")
            hidden_sheets = set()

        # 从XLSX内部XML读取合并单元格信息（无需openpyxl完整加载）
        merged_cells_map = self._get_merged_cells_from_xlsx(file_path)

        # 使用pandas读取可见Sheet的数据
        xls = pd.ExcelFile(file_path)
        sheet_names = xls.sheet_names

        for sheet_name in sheet_names:
            # 跳过隐藏的Sheet
            if sheet_name in hidden_sheets:
                continue

            try:
                df = pd.read_excel(xls, sheet_name=sheet_name, header=None)
            except Exception as e:
                print(f"警告：无法读取Sheet '{sheet_name}': {e}")
                continue

            # 填充合并单元格的值（与sheet加载分离，失败不影响sheet可用）
            if sheet_name in merged_cells_map:
                try:
                    df = df.astype(object)
                    for (r, c), (tl_r, tl_c) in merged_cells_map[sheet_name].items():
                        if r < len(df) and c < len(df.columns):
                            df.iloc[r, c] = df.iloc[tl_r, tl_c]
                except Exception:
                    pass

            workbook.sheets[sheet_name] = df

        return workbook

    def validate_structure(self, workbook: AuditWorkbook) -> list:
        """验证底稿结构是否符合标准"""
        missing_sheets = []
        for standard_sheet in self.STANDARD_SHEETS:
            # 模糊匹配，允许尾部空格等差异
            found = False
            for actual_sheet in workbook.sheets:
                if actual_sheet.strip() == standard_sheet.strip():
                    found = True
                    break
            if not found:
                missing_sheets.append(standard_sheet)
        return missing_sheets

    def get_sheet_summary(self, workbook: AuditWorkbook) -> dict:
        """获取每个Sheet的概要信息"""
        summary = {}
        for sheet_name, df in workbook.sheets.items():
            summary[sheet_name] = {
                "rows": len(df),
                "cols": len(df.columns),
                "has_data": not df.empty
            }
        return summary

    def find_cell_by_keyword(self, workbook: AuditWorkbook,
                              sheet_name: str, keyword: str,
                              search_range: Optional[Tuple[int, int]] = None) -> Optional[Tuple[int, int]]:
        """
        通过关键词动态搜索单元格位置

        Args:
            workbook: 工作簿对象
            sheet_name: Sheet名称（支持模糊匹配）
            keyword: 搜索关键词
            search_range: (起始行, 结束行) 搜索范围，None则搜索全部

        Returns:
            (row, col) 或 None（未找到）
        """
        # 模糊匹配Sheet名称
        actual_sheet = self._match_sheet_name(workbook, sheet_name)
        if actual_sheet is None:
            return None

        df = workbook.sheets[actual_sheet]
        start_row = search_range[0] if search_range else 0
        end_row = search_range[1] if search_range else len(df)

        for row_idx in range(start_row, min(end_row, len(df))):
            for col_idx in range(len(df.columns)):
                value = df.iloc[row_idx, col_idx]
                if pd.notna(value) and keyword in str(value).strip():
                    self._record_evidence(actual_sheet, row_idx, col_idx)
                    return (row_idx, col_idx)

        return None

    def find_all_cells_by_keyword(self, workbook: AuditWorkbook,
                                   sheet_name: str, keyword: str,
                                   search_range: Optional[Tuple[int, int]] = None) -> List[Tuple[int, int]]:
        """
        通过关键词搜索所有匹配的单元格位置

        Returns:
            [(row, col), ...] 匹配的单元格列表
        """
        actual_sheet = self._match_sheet_name(workbook, sheet_name)
        if actual_sheet is None:
            return []

        df = workbook.sheets[actual_sheet]
        start_row = search_range[0] if search_range else 0
        end_row = search_range[1] if search_range else len(df)

        results = []
        for row_idx in range(start_row, min(end_row, len(df))):
            for col_idx in range(len(df.columns)):
                value = df.iloc[row_idx, col_idx]
                if pd.notna(value) and keyword in str(value).strip():
                    self._record_evidence(actual_sheet, row_idx, col_idx)
                    results.append((row_idx, col_idx))

        return results

    def find_row_by_keyword(self, workbook: AuditWorkbook,
                             sheet_name: str, keyword: str,
                             search_range: Optional[Tuple[int, int]] = None) -> Optional[int]:
        """
        通过关键词搜索行号（找到关键词所在的行）

        Returns:
            row_index 或 None
        """
        cell = self.find_cell_by_keyword(workbook, sheet_name, keyword, search_range)
        return cell[0] if cell else None

    def get_cell_value(self, workbook: AuditWorkbook,
                        sheet_name: str, row: int, col: int) -> Optional[object]:
        """
        获取指定单元格的值（动态定位后使用）

        Args:
            row, col: 0-based索引

        Returns:
            单元格值 或 None
        """
        actual_sheet = self._match_sheet_name(workbook, sheet_name)
        if actual_sheet is None:
            return None

        df = workbook.sheets[actual_sheet]
        if row < len(df) and col < len(df.columns):
            self._record_evidence(actual_sheet, row, col)
            return df.iloc[row, col]
        return None

    def get_row_data(self, workbook: AuditWorkbook,
                      sheet_name: str, row: int,
                      start_col: int = 0, end_col: Optional[int] = None) -> List[object]:
        """
        获取指定行的数据（从start_col到end_col）

        Args:
            row: 0-based行号
            start_col: 起始列（0-based）
            end_col: 结束列（None则取到最后一列）

        Returns:
            该行的值列表
        """
        actual_sheet = self._match_sheet_name(workbook, sheet_name)
        if actual_sheet is None:
            return []

        df = workbook.sheets[actual_sheet]
        if row >= len(df):
            return []

        if end_col is None:
            end_col = len(df.columns)

        values = []
        for col_idx in range(start_col, min(end_col, len(df.columns))):
            self._record_evidence(actual_sheet, row, col_idx)
            values.append(df.iloc[row, col_idx])

        return values

    def find_value_in_row(self, workbook: AuditWorkbook,
                           sheet_name: str, keyword: str,
                           value_col_offset: int = 1,
                           search_range: Optional[Tuple[int, int]] = None) -> Optional[object]:
        """
        找到关键词所在行，然后获取其右侧偏移列的值

        例如：搜索"年初余额"，然后取同一行第4列（偏移3列）的数值

        Args:
            keyword: 搜索关键词
            value_col_offset: 值所在列相对于关键词列的偏移量

        Returns:
            目标值 或 None
        """
        cell = self.find_cell_by_keyword(workbook, sheet_name, keyword, search_range)
        if cell is None:
            return None

        row, col = cell
        value_col = col + value_col_offset
        return self.get_cell_value(workbook, sheet_name, row, value_col)

    def _match_sheet_name(self, workbook: AuditWorkbook, target_name: str) -> Optional[str]:
        """
        模糊匹配Sheet名称（处理尾部空格等差异）

        Args:
            target_name: 目标Sheet名称

        Returns:
            实际Sheet名称 或 None
        """
        # 先精确匹配
        if target_name in workbook.sheets:
            return target_name

        # 再模糊匹配（去除尾部空格）
        target_stripped = target_name.strip()
        for actual_name in workbook.sheets:
            if actual_name.strip() == target_stripped:
                return actual_name

        # 最后部分匹配（包含关系）
        for actual_name in workbook.sheets:
            if target_stripped in actual_name.strip() or actual_name.strip() in target_stripped:
                return actual_name

        return None

    def find_sheets_by_keyword(self, workbook: AuditWorkbook, keyword: str) -> List[str]:
        """
        根据关键词模糊匹配所有Sheet名称

        Args:
            workbook: 工作簿对象
            keyword: 搜索关键词（如"02.1a"、"SAP"、"by item"等）

        Returns:
            匹配的所有Sheet名称列表
        """
        results = []
        keyword_lower = keyword.lower()
        for sheet_name in workbook.sheets:
            if keyword_lower in sheet_name.lower():
                self._record_evidence(sheet_name, text=f"匹配Sheet：{sheet_name}")
                results.append(sheet_name)
        return results

    def identify_depreciation_method(self, workbook: AuditWorkbook, sheet_name: str = None) -> str:
        """
        识别折旧测试方法：SAP / TOD / By Item

        优先级：sheet名称 → 表格内容结构
        - 含"by item"/"by_item" → 直接判By_Item
        - 含"SAP"（且不含by item）→ 直接判SAP
        - 仅含"TOD"或无法从名称判断 → 走表格内容判断

        Args:
            workbook: 工作簿对象
            sheet_name: Sheet名称（可选，不指定则自动查找K.03.2相关sheet）

        Returns:
            "SAP" / "TOD" / "By_Item" / "Unknown"
        """
        # 查找目标Sheet
        target_sheets = []
        if sheet_name:
            actual = self._match_sheet_name(workbook, sheet_name)
            if actual:
                target_sheets = [actual]
        else:
            target_sheets = self.find_sheets_by_keyword(workbook, "K.03.2")

        if not target_sheets:
            # 尝试模糊匹配含"折旧"字样的sheet
            target_sheets = self.find_sheets_by_keyword(workbook, "折旧")

        for actual_sheet in target_sheets:
            df = workbook.sheets[actual_sheet]
            if df is None or df.empty:
                continue

            # === 第一层：sheet名称判断 ===
            name_lower = actual_sheet.lower().replace(" ", "").replace("_", "")

            # 含"byitem" → 一定是By_Item，不看内容
            if "byitem" in name_lower:
                return "By_Item"

            # 含"sap"（且不含byitem）→ 一定是SAP
            if "sap" in name_lower:
                return "SAP"

            # === 第二层：表格内容判断（TOD vs By_Item）===
            # 策略1：检测By Item特征（固定资产编号+累计折旧+原值的列结构）
            byitem_score = 0
            # By Item特征列（Row 8附近）
            byitem_keywords = [
                ("固定资产编号", [2]),       # C列
                ("原值", [7, 8]),           # H/I列
                ("累计折旧", [8, 9]),        # I/J列
                ("本期差异", [17]),          # R列
                ("本期应折旧", [16]),        # Q列
                ("使用寿命", [5]),           # F列
                ("入账开始日期", [4]),        # E列
            ]
            for kw, expected_cols in byitem_keywords:
                for row_idx in range(0, min(15, len(df))):
                    for col_idx in range(len(df.columns)):
                        val = df.iloc[row_idx, col_idx]
                        if pd.notna(val) and kw in str(val).strip():
                            self._record_evidence(actual_sheet, row_idx, col_idx)
                            if col_idx in expected_cols:
                                byitem_score += 2
                            else:
                                byitem_score += 0.5

            # 策略2：检测TOD特征（样本数量+样本类型+差异列）
            tod_score = 0
            # TOD特征列（Row 31附近）
            tod_keywords = [
                ("样本数量", [1]),            # B列
                ("样本类型", [2]),            # C列
                ("固定资产类别", [3]),         # D列
                ("差异", [15]),              # P列
                ("获得的证据", [16]),         # Q列
                ("EY Works", [17]),          # R列（标记列）
            ]
            for kw, expected_cols in tod_keywords:
                for row_idx in range(min(25, len(df)), min(40, len(df))):
                    for col_idx in range(len(df.columns)):
                        val = df.iloc[row_idx, col_idx]
                        if pd.notna(val) and kw in str(val).strip():
                            self._record_evidence(actual_sheet, row_idx, col_idx)
                            if col_idx in expected_cols:
                                tod_score += 2
                            else:
                                tod_score += 0.5

            # 如果有明确的得分差异，返回得分较高的方法
            if byitem_score >= 4 and byitem_score > tod_score:
                return "By_Item"
            if tod_score >= 4 and tod_score > byitem_score:
                return "TOD"

        return "Unknown"

    def get_depreciation_sheet_info(self, workbook: AuditWorkbook) -> dict:
        """
        获取底稿中所有折旧测试Sheet的信息，包括识别出的测试方法

        Returns:
            {
                "sheets": [sheet_name, ...],
                "methods": {sheet_name: "SAP"|"TOD"|"By_Item"|"Unknown", ...}
            }
        """
        info = {"sheets": [], "methods": {}}

        # 收集所有折旧相关Sheet（去重）
        all_sheets = set()

        # SAP相关：K.03.1 / SAP
        sap_sheets = self.find_sheets_by_keyword(workbook, "K.03.1")
        if not sap_sheets:
            sap_sheets = self.find_sheets_by_keyword(workbook, "SAP")
        for s in sap_sheets:
            all_sheets.add(s)

        # TOD/By_Item相关：K.03.2 / 折旧测试
        k032_sheets = self.find_sheets_by_keyword(workbook, "K.03.2")
        if not k032_sheets:
            k032_sheets = self.find_sheets_by_keyword(workbook, "折旧")
        for s in k032_sheets:
            all_sheets.add(s)

        for sheet in all_sheets:
            method = self.identify_depreciation_method(workbook, sheet)
            info["sheets"].append(sheet)
            info["methods"][sheet] = method

        return info
